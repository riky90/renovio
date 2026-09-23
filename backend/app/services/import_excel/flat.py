"""Lettore generico del formato piatto d'export (D-01, D-03, D-09, D-12).

Ogni foglio dati scritto da `export_excel/workbook.py::scrivi_foglio_dati`
ha `id` come prima colonna dell'header ed e' accompagnato dal foglio
`ISTRUZIONI`. Questo modulo fa il percorso inverso: riconosce il formato,
legge l'header, e trasforma ogni riga in un `RigaPiatta` che non sa niente
di modelli o database — legge bytes e restituisce righe, nessuna scrittura.

Regola cardine (CUR-02, D-09): una cella VUOTA nel foglio dati significa
"campo non toccato" e la sua chiave non entra in `RigaPiatta.valori`. Una
cella che contiene `TOKEN_CANCELLA` (riesportato da `export_excel.workbook`,
mai ridefinito qui: una sola definizione del token in tutto il progetto)
entra invece in `valori` con valore `None` — e' l'unico modo per svuotare
un campo gia' valorizzato a database. La distinzione fra "chiave assente" e
"chiave con valore None" resta intatta fino all'upsert
(`listino_upload_service.py`), che e' l'unico punto che conosce quali
colonne sono nullable a database e puo' quindi decidere se un `None`
esplicito e' applicabile o va scartato.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime

from openpyxl.workbook import Workbook

from app.schemas.admin_upload import ScartoRowOut
from app.services.export_excel.workbook import TOKEN_CANCELLA

FOGLIO_ISTRUZIONI = "ISTRUZIONI"
_CELLA_TIMESTAMP = "B4"
_COLONNA_ID = "id"


@dataclass
class RigaPiatta:
    """Una riga dati del foglio piatto. `valori` contiene SOLO le colonne
    che questa riga porta con se': assenza = campo non toccato (D-09).

    `id` e' intero per le risorse a PK intera (comportamento storico, 08-03).
    Il piano 08-06 introduce `id_numerico=False` in `leggi_foglio_piatto`
    per le risorse a PK testuale (es. `Extra.code`): in quel caso `id` resta
    la stringa grezza, mai forzata a `int`."""

    numero_riga: int
    id: int | str | None
    valori: dict[str, object] = field(default_factory=dict)


def _cella_vuota(valore: object) -> bool:
    return valore is None or (isinstance(valore, str) and valore.strip() == "")


# Testi accettati come booleano in una cella scritta a mano. Excel non
# impone il tipo booleano a una cella: l'admin che vuole disattivare una
# riga ci scrive quello che gli viene, in italiano o in inglese.
_TESTI_VERI = {"true", "vero", "1", "si", "sì", "s", "x", "y", "yes"}
_TESTI_FALSI = {"false", "falso", "0", "no", "n"}


def valore_booleano(valore: object) -> bool | None:
    """Normalizza una cella in un booleano, o `None` se non e' decidibile.

    `bool(valore)` non basta e non e' innocuo: in Python `bool("FALSE")`,
    `bool("no")` e `bool("0")` sono tutti `True`. Un admin che scrive
    `FALSE` a mano in una cella `attivo` — cosa che Excel permette, perche'
    la cella non ha un tipo booleano imposto — vedrebbe la riga restare
    attiva, con l'intenzione esattamente opposta e nessun errore da nessuna
    parte.

    Il round-trip a vuoto non cambia: l'export scrive il booleano nativo e
    openpyxl lo rilegge come `bool`, che passa per il ramo diretto qui
    sotto.

    `None` significa "non so", e chi chiama lo traduce in uno scarto
    descritto — mai in un'ipotesi silenziosa su cosa intendesse l'admin.
    """
    if isinstance(valore, bool):
        return valore
    if isinstance(valore, (int, float)):
        if valore == 1:
            return True
        if valore == 0:
            return False
        return None
    if isinstance(valore, str):
        testo = valore.strip().lower()
        if testo in _TESTI_VERI:
            return True
        if testo in _TESTI_FALSI:
            return False
    return None


def e_formato_piatto(wb: Workbook) -> bool:
    """Discriminatore di D-01: il formato piatto ha SEMPRE il foglio
    `ISTRUZIONI` e almeno un foglio dati la cui prima colonna di header e'
    esattamente `id`. Il computo originale (foglio `A-EDILE` + cataloghi
    B/C/D) non ha mai un foglio `ISTRUZIONI` -> restituisce sempre False."""
    if FOGLIO_ISTRUZIONI not in wb.sheetnames:
        return False
    for nome in wb.sheetnames:
        if nome == FOGLIO_ISTRUZIONI:
            continue
        prima_cella = wb[nome].cell(row=1, column=1).value
        if isinstance(prima_cella, str) and prima_cella.strip() == _COLONNA_ID:
            return True
    return False


def leggi_timestamp_export(wb: Workbook) -> datetime | None:
    """Legge il timestamp di generazione scritto da `nuovo_workbook` in
    `ISTRUZIONI!B4`. Restituisce `None` se il foglio o la cella sono
    assenti/illeggibili — nessuna eccezione, il campo e' informativo
    (consumato da D-11 nel piano 08-08, qui solo esposto)."""
    if FOGLIO_ISTRUZIONI not in wb.sheetnames:
        return None
    valore = wb[FOGLIO_ISTRUZIONI][_CELLA_TIMESTAMP].value
    if not isinstance(valore, str):
        return None
    try:
        return datetime.fromisoformat(valore)
    except ValueError:
        return None


def foglio_mancante(wb: Workbook, nome: str) -> ScartoRowOut | None:
    """Un foglio dati atteso e assente produce uno scarto descritto, mai
    un'eccezione che azzera l'intero import."""
    if nome in wb.sheetnames:
        return None
    return ScartoRowOut(riga=0, motivo=f"foglio '{nome}' mancante")


def leggi_foglio_piatto(
    wb: Workbook, nome_foglio: str, *, colonne_ammesse: set[str], id_numerico: bool = True
) -> tuple[list[RigaPiatta], list[ScartoRowOut]]:
    """Legge l'header dalla prima riga del foglio `nome_foglio` e produce
    una `RigaPiatta` per ogni riga dati non completamente vuota.

    Le colonne dell'header che non sono in `colonne_ammesse` vengono
    ignorate silenziosamente (un admin che aggiunge una colonna di appunti
    non deve far fallire l'import). Una colonna di `colonne_ammesse`
    assente dall'header semplicemente non compare mai in nessuna
    `RigaPiatta.valori`: non e' un errore, e' esattamente il caso di CUR-02.

    `id_numerico=False` (piano 08-06): l'`id` resta la stringa grezza,
    senza tentare `int(...)` — necessario per le risorse a PK testuale
    (`Extra.code`), dove forzare un intero scarterebbe SEMPRE la riga.

    Scarti (con motivo esplicito, mai un'eccezione): riga completamente
    vuota, riga il cui `id` non e' un intero (solo quando `id_numerico`).
    """
    scarti: list[ScartoRowOut] = []
    ws = wb[nome_foglio]

    prima_riga = next(ws.iter_rows(min_row=1, max_row=1), ())
    header: list[str | None] = [
        cella.value.strip() if isinstance(cella.value, str) else None for cella in prima_riga
    ]

    righe: list[RigaPiatta] = []
    for numero_riga, row in enumerate(ws.iter_rows(min_row=2), start=2):
        grezzo: dict[str, object] = {}
        for indice, cella in enumerate(row):
            if indice >= len(header):
                continue
            nome_colonna = header[indice]
            if nome_colonna is None or nome_colonna not in colonne_ammesse:
                continue
            grezzo[nome_colonna] = cella.value

        if all(_cella_vuota(v) for v in grezzo.values()):
            scarti.append(ScartoRowOut(riga=numero_riga, motivo="riga vuota"))
            continue

        id_raw = grezzo.get(_COLONNA_ID)
        riga_id: int | str | None = None
        if not _cella_vuota(id_raw):
            if id_numerico:
                try:
                    riga_id = int(id_raw)  # type: ignore[arg-type]
                except (TypeError, ValueError):
                    scarti.append(ScartoRowOut(riga=numero_riga, motivo=f"id non numerico: {id_raw!r}"))
                    continue
            else:
                riga_id = str(id_raw).strip()
                if not riga_id:
                    scarti.append(ScartoRowOut(riga=numero_riga, motivo="id vuoto"))
                    continue

        valori: dict[str, object] = {}
        for nome_colonna, valore in grezzo.items():
            if nome_colonna == _COLONNA_ID:
                continue
            if _cella_vuota(valore):
                continue
            if isinstance(valore, str) and valore.strip() == TOKEN_CANCELLA:
                valori[nome_colonna] = None
            else:
                valori[nome_colonna] = valore

        righe.append(RigaPiatta(numero_riga=numero_riga, id=riga_id, valori=valori))

    return righe, scarti
