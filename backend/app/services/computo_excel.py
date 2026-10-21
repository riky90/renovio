"""Excel del computo parametrico (Fase 11, CUR3-01): esporta e reimporta
grandezze, termini, regole, valori delle opzioni e preset in un unico file.

E' la lingua del tecnico: il computo 2025 aveva gia' cominciato a scrivere le
regole nelle colonne T/U, qui le stesse informazioni stanno in colonne
strutturate (tipo di regola, grandezza, fattore, sorgenti...).

L'import SOSTITUISCE in blocco le sette tabelle del computo: il file e'
l'intero ragionamento, non una patch riga per riga. Per questo e' tutto o
niente: si scrive dentro un SAVEPOINT, si valida con `valida_computo` e, se
c'e' anche un solo errore (o se e' un'anteprima), si torna indietro. Un file
sbagliato non lascia mai il computo a meta'.

Nessun coefficiente impresa: il file porta quantita', fattori e flag, e della
voce solo codice, descrizione e unita' di misura come promemoria in sola
lettura.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.models.catalogo import TipoRistrutturazione
from app.models.computo import (
    Grandezza,
    PresetRisposta,
    RegolaVoce,
    ScaglioneRegola,
    SorgenteRegola,
    TermineGrandezza,
    ValoreGrandezzaOpzione,
)
from app.models.dettaglio_pro import Voce
from app.models.domanda import OpzioneDomanda
from app.services.computo import TIPI_REGOLA, valida_computo

FOGLIO_ISTRUZIONI = "ISTRUZIONI"

COLONNE: dict[str, list[str]] = {
    "grandezze": [
        "code", "nome", "um", "descrizione", "ordine", "costante", "da_attesa",
        "arrotondamento", "minimo", "massimo", "attiva",
    ],
    "termini": ["grandezza", "sorgente", "fattore", "per_grandezza"],
    "regole": [
        "voce", "descrizione_voce", "um_voce", "tipo", "grandezza", "fattore",
        "condizione", "importo_minimo", "sorgenti", "scaglioni", "nota", "attiva",
    ],
    "valori_opzione": ["domanda", "opzione", "grandezza", "valore"],
    "preset": ["tipo", "domanda", "opzione", "valore"],
}

#: Colonne scritte come promemoria e ignorate all'import.
SOLO_LETTURA = {"descrizione_voce", "um_voce"}

ISTRUZIONI = [
    "Computo parametrico — grandezze della casa, regole per voce, preset per tipo.",
    "",
    "L'import SOSTITUISCE tutto il computo con il contenuto di questo file, ed e' tutto o niente:",
    "se c'e' un errore non cambia nulla e l'anteprima elenca foglio, riga e motivo.",
    "",
    "grandezze: default = costante + somma dei termini; da_attesa = prendi la fascia di attesa_immobile.",
    "termini: fattore x sorgente [x per_grandezza]. Sorgenti: 'mq', 'tipo_<codice>' o un'altra grandezza.",
    f"regole: tipo fra {', '.join(sorted(TIPI_REGOLA))}.",
    "  sorgenti = codici di voce o sottocategoria separati da virgola (da_voci, percentuale...).",
    "  scaglioni = 'fino_a:valore' separati da ';' — 'oltre:valore' per l'ultimo (es. 100:1; 150:1.25; oltre:1.5).",
    "  condizione = la voce entra solo se quella grandezza e' > 0.",
    "descrizione_voce e um_voce sono un promemoria: all'import si ignorano.",
    "valori_opzione: l'opzione di una domanda assegna un valore a una grandezza.",
    "preset: risposta di default di una domanda per un tipo di ristrutturazione.",
]


def _testo(valore: Decimal | None) -> Any:
    if valore is None:
        return None
    normalizzato = valore.normalize()
    return int(normalizzato) if normalizzato == normalizzato.to_integral_value() else float(normalizzato)


def esporta_computo(session: Session) -> tuple[bytes, int]:
    """Il workbook del computo e il numero di righe dati scritte."""
    wb = Workbook()
    istruzioni = wb.active
    istruzioni.title = FOGLIO_ISTRUZIONI
    for riga in ISTRUZIONI:
        istruzioni.append([riga])
    istruzioni.column_dimensions["A"].width = 120

    righe: dict[str, list[list[Any]]] = {nome: [] for nome in COLONNE}

    for g in session.query(Grandezza).order_by(Grandezza.ordine, Grandezza.code).all():
        righe["grandezze"].append([
            g.code, g.nome, g.um, g.descrizione, g.ordine, _testo(g.costante), g.da_attesa,
            g.arrotondamento, _testo(g.minimo), _testo(g.massimo), g.attiva,
        ])
    for t in session.query(TermineGrandezza).order_by(
        TermineGrandezza.grandezza_code, TermineGrandezza.ordine, TermineGrandezza.id
    ).all():
        righe["termini"].append([t.grandezza_code, t.sorgente_code, _testo(t.fattore), t.per_grandezza_code])

    voci = {v.codice: v for v in session.query(Voce).filter(Voce.codice.isnot(None)).all()}
    sorgenti: dict[str, list[str]] = {}
    for s in session.query(SorgenteRegola).order_by(SorgenteRegola.id).all():
        sorgenti.setdefault(s.voce_codice, []).append(s.sorgente_codice)
    scaglioni: dict[str, list[ScaglioneRegola]] = {}
    for s in session.query(ScaglioneRegola).all():
        scaglioni.setdefault(s.voce_codice, []).append(s)
    for r in session.query(RegolaVoce).order_by(RegolaVoce.voce_codice).all():
        voce = voci.get(r.voce_codice)
        testo_scaglioni = "; ".join(
            f"{'oltre' if s.fino_a is None else _testo(s.fino_a)}:{_testo(s.valore)}"
            for s in sorted(
                scaglioni.get(r.voce_codice, []),
                key=lambda s: (s.fino_a is None, s.fino_a or Decimal("0")),
            )
        )
        righe["regole"].append([
            r.voce_codice,
            voce.descrizione[:120] if voce else None,
            voce.um if voce else None,
            r.tipo_regola,
            r.grandezza_code,
            _testo(r.fattore),
            r.condizione_grandezza_code,
            _testo(r.importo_minimo),
            ", ".join(sorgenti.get(r.voce_codice, [])) or None,
            testo_scaglioni or None,
            r.nota,
            r.attiva,
        ])

    for valore, opzione in (
        session.query(ValoreGrandezzaOpzione, OpzioneDomanda)
        .join(OpzioneDomanda, OpzioneDomanda.id == ValoreGrandezzaOpzione.opzione_id)
        .order_by(OpzioneDomanda.domanda_code, OpzioneDomanda.ordine, ValoreGrandezzaOpzione.grandezza_code)
        .all()
    ):
        righe["valori_opzione"].append([opzione.domanda_code, opzione.code, valore.grandezza_code, _testo(valore.valore)])

    for p in session.query(PresetRisposta).order_by(
        PresetRisposta.tipo_code, PresetRisposta.domanda_code, PresetRisposta.id
    ).all():
        righe["preset"].append([p.tipo_code, p.domanda_code, p.opzione_code, _testo(p.valore)])

    totale = 0
    for nome, colonne in COLONNE.items():
        ws = wb.create_sheet(nome)
        ws.append(colonne)
        for cella in ws[1]:
            cella.font = Font(bold=True)
        for riga in righe[nome]:
            ws.append(riga)
        totale += len(righe[nome])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), totale


# --- import -----------------------------------------------------------------


@dataclass
class EsitoImportComputo:
    errori: list[str] = field(default_factory=list)
    conteggi: dict[str, int] = field(default_factory=dict)
    applicato: bool = False


class _Riga:
    def __init__(self, foglio: str, numero: int, valori: dict[str, Any], errori: list[str]):
        self.foglio, self.numero, self.valori, self._errori = foglio, numero, valori, errori

    def errore(self, messaggio: str) -> None:
        self._errori.append(f"{self.foglio}, riga {self.numero}: {messaggio}")

    def testo(self, colonna: str, obbligatorio: bool = False) -> str | None:
        valore = self.valori.get(colonna)
        if valore is None or (isinstance(valore, str) and not valore.strip()):
            if obbligatorio:
                self.errore(f"'{colonna}' e' obbligatorio.")
            return None
        return str(valore).strip()

    def numero_dec(self, colonna: str, obbligatorio: bool = False, default: Decimal | None = None) -> Decimal | None:
        valore = self.valori.get(colonna)
        if valore is None or (isinstance(valore, str) and not valore.strip()):
            if obbligatorio and default is None:
                self.errore(f"'{colonna}' e' obbligatorio.")
            return default
        if isinstance(valore, bool):
            self.errore(f"'{colonna}' deve essere un numero.")
            return default
        try:
            return Decimal(str(valore).strip().replace(",", "."))
        except InvalidOperation:
            self.errore(f"'{colonna}' deve essere un numero (trovato {valore!r}).")
            return default

    def booleano(self, colonna: str, default: bool) -> bool:
        valore = self.valori.get(colonna)
        if valore is None or (isinstance(valore, str) and not valore.strip()):
            return default
        if isinstance(valore, bool):
            return valore
        testo = str(valore).strip().lower()
        if testo in {"true", "vero", "1", "si", "sì", "s", "x", "yes"}:
            return True
        if testo in {"false", "falso", "0", "no", "n"}:
            return False
        self.errore(f"'{colonna}' deve essere si'/no (trovato {valore!r}).")
        return default


def _leggi_fogli(wb, errori: list[str]) -> dict[str, list[_Riga]]:
    fogli: dict[str, list[_Riga]] = {}
    for nome, colonne in COLONNE.items():
        if nome not in wb.sheetnames:
            errori.append(f"Manca il foglio '{nome}'.")
            continue
        ws = wb[nome]
        intestazione = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        mancanti = [c for c in colonne if c not in intestazione and c not in SOLO_LETTURA]
        if mancanti:
            errori.append(f"{nome}: mancano le colonne {', '.join(mancanti)}.")
            continue
        righe = []
        for numero, valori in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in valori):
                continue
            mappa = {intestazione[i]: v for i, v in enumerate(valori) if i < len(intestazione)}
            righe.append(_Riga(nome, numero, mappa, errori))
        fogli[nome] = righe
    return fogli


def importa_computo(session: Session, wb, *, applica: bool) -> EsitoImportComputo:
    """Sostituisce il computo con il contenuto del workbook. Con
    `applica=False` (anteprima) o con errori non cambia nulla."""
    esito = EsitoImportComputo()
    fogli = _leggi_fogli(wb, esito.errori)
    if esito.errori:
        return esito

    savepoint = session.begin_nested()
    try:
        for modello in (
            PresetRisposta,
            ValoreGrandezzaOpzione,
            ScaglioneRegola,
            SorgenteRegola,
            RegolaVoce,
            TermineGrandezza,
            Grandezza,
        ):
            session.query(modello).delete(synchronize_session=False)
        session.flush()

        codici_grandezza: set[str] = set()
        for r in fogli["grandezze"]:
            code = r.testo("code", obbligatorio=True)
            nome = r.testo("nome", obbligatorio=True)
            if code is None or nome is None:
                continue
            if code in codici_grandezza:
                r.errore(f"grandezza {code!r} ripetuta.")
                continue
            codici_grandezza.add(code)
            session.add(
                Grandezza(
                    code=code,
                    nome=nome,
                    um=r.testo("um") or "",
                    descrizione=r.testo("descrizione"),
                    ordine=int(r.numero_dec("ordine", default=Decimal("0")) or 0),
                    costante=r.numero_dec("costante", default=Decimal("0")),
                    da_attesa=r.booleano("da_attesa", False),
                    arrotondamento=r.testo("arrotondamento") or "nessuno",
                    minimo=r.numero_dec("minimo"),
                    massimo=r.numero_dec("massimo"),
                    attiva=r.booleano("attiva", True),
                )
            )
        session.flush()

        for ordine, r in enumerate(fogli["termini"]):
            grandezza = r.testo("grandezza", obbligatorio=True)
            sorgente = r.testo("sorgente", obbligatorio=True)
            fattore = r.numero_dec("fattore", obbligatorio=True)
            if grandezza is None or sorgente is None or fattore is None:
                continue
            if grandezza not in codici_grandezza:
                r.errore(f"grandezza {grandezza!r} non definita nel foglio grandezze.")
                continue
            session.add(
                TermineGrandezza(
                    grandezza_code=grandezza,
                    sorgente_code=sorgente,
                    fattore=fattore,
                    per_grandezza_code=r.testo("per_grandezza"),
                    ordine=ordine,
                )
            )

        codici_regola: set[str] = set()
        for r in fogli["regole"]:
            voce = r.testo("voce", obbligatorio=True)
            tipo = r.testo("tipo", obbligatorio=True)
            if voce is None or tipo is None:
                continue
            if voce in codici_regola:
                r.errore(f"regola per la voce {voce!r} ripetuta.")
                continue
            codici_regola.add(voce)
            session.add(
                RegolaVoce(
                    voce_codice=voce,
                    tipo_regola=tipo,
                    grandezza_code=r.testo("grandezza"),
                    fattore=r.numero_dec("fattore", default=Decimal("1")),
                    condizione_grandezza_code=r.testo("condizione"),
                    importo_minimo=r.numero_dec("importo_minimo"),
                    nota=r.testo("nota"),
                    attiva=r.booleano("attiva", True),
                )
            )
            session.flush()
            for codice in (r.testo("sorgenti") or "").split(","):
                if codice.strip():
                    session.add(SorgenteRegola(voce_codice=voce, sorgente_codice=codice.strip()))
            for pezzo in (r.testo("scaglioni") or "").split(";"):
                if not pezzo.strip():
                    continue
                if ":" not in pezzo:
                    r.errore(f"scaglione {pezzo.strip()!r}: usa 'fino_a:valore'.")
                    continue
                soglia, valore = (p.strip() for p in pezzo.split(":", 1))
                try:
                    fino_a = None if soglia.lower() == "oltre" else Decimal(soglia.replace(",", "."))
                    session.add(
                        ScaglioneRegola(
                            voce_codice=voce, fino_a=fino_a, valore=Decimal(valore.replace(",", "."))
                        )
                    )
                except InvalidOperation:
                    r.errore(f"scaglione {pezzo.strip()!r}: numeri non validi.")

        for r in fogli["valori_opzione"]:
            domanda = r.testo("domanda", obbligatorio=True)
            codice_opzione = r.testo("opzione", obbligatorio=True)
            grandezza = r.testo("grandezza", obbligatorio=True)
            valore = r.numero_dec("valore", obbligatorio=True)
            if None in (domanda, codice_opzione, grandezza, valore):
                continue
            opzione = (
                session.query(OpzioneDomanda)
                .filter_by(domanda_code=domanda, code=codice_opzione)
                .one_or_none()
            )
            if opzione is None:
                r.errore(f"opzione {codice_opzione!r} della domanda {domanda!r} inesistente.")
                continue
            if grandezza not in codici_grandezza:
                r.errore(f"grandezza {grandezza!r} non definita nel foglio grandezze.")
                continue
            session.add(ValoreGrandezzaOpzione(opzione_id=opzione.id, grandezza_code=grandezza, valore=valore))

        for r in fogli["preset"]:
            tipo = r.testo("tipo", obbligatorio=True)
            domanda = r.testo("domanda", obbligatorio=True)
            if tipo is None or domanda is None:
                continue
            if session.get(TipoRistrutturazione, tipo) is None:
                r.errore(f"tipo di ristrutturazione {tipo!r} inesistente.")
                continue
            opzione = r.testo("opzione")
            valore = r.numero_dec("valore")
            if opzione is None and valore is None:
                r.errore("serve un'opzione o un valore.")
                continue
            if opzione is not None and (
                session.query(OpzioneDomanda).filter_by(domanda_code=domanda, code=opzione).one_or_none()
                is None
            ):
                r.errore(f"opzione {opzione!r} della domanda {domanda!r} inesistente.")
                continue
            session.add(PresetRisposta(tipo_code=tipo, domanda_code=domanda, opzione_code=opzione, valore=valore))

        session.flush()
        if not esito.errori:
            esito.errori.extend(valida_computo(session))
            # Una voce del listino senza codice non puo' avere una regola, ma
            # una regola su un codice senza voce si': `valida_computo` lo dice.
        esito.conteggi = {nome: len(righe) for nome, righe in fogli.items()}

        if esito.errori or not applica:
            savepoint.rollback()
            return esito
        savepoint.commit()
        session.commit()
        esito.applicato = True
        return esito
    except Exception:
        savepoint.rollback()
        raise
