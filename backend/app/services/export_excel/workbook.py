"""Shell comune a tutti gli export Excel del pannello admin (D-09, D-10, D-11,
D-13). Ogni file esportato da questa fase (listino, domande, parametri, lead)
passa da qui per il foglio ISTRUZIONI, il token di cancellazione, il naming e
la serializzazione — cosi' le quattro regole restano identiche ovunque
invece di essere riscritte quattro volte.

Rappresentazione dei `Decimal` (round-trip a zero, D-12/08-03): ogni valore
`Decimal` viene scritto come `float`. I prezzi di questo dominio hanno al
massimo 2 cifre decimali (`Numeric(12, 2)`), e un `float` IEEE-754 rappresenta
un numero con 2 decimali con un errore ben sotto la precisione che Excel
mostra o che l'utente puo' modificare a mano; rileggendo con openpyxl si
ottiene un `float` che, ri-arrotondato a 2 decimali e passato a
`Decimal(str(...))`, coincide col valore originale. La stringa e' stata
scartata deliberatamente: avrebbe reso le celle testo invece di numeri,
rompendo la modifica a mano in Excel (funzioni, formattazione valuta) che e'
il punto stesso del file scaricabile.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import Workbook

TOKEN_CANCELLA = "#CANCELLA#"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TESTO_REGOLA_CELLA_VUOTA = (
    "Cella vuota = campo non modificato. Per cancellare un valore esistente "
    f"scrivi {TOKEN_CANCELLA} nella cella."
)
_TESTO_NO_ELIMINAZIONE = (
    "Le righe non si eliminano da questo file: l'eliminazione passa dal pannello."
)
_TESTO_NO_TOCCARE_ID = "Non modificare la colonna id: e' la chiave con cui il file viene ricaricato."
_TESTO_AVVISO_MARGINE = (
    "Questo file contiene il coefficiente impresa (margine). Uso riservato al "
    "ruolo admin. Non inoltrare."
)


def nuovo_workbook(*, con_margine: bool = False) -> Workbook:
    """Crea il workbook con il foglio `ISTRUZIONI` come primo foglio (D-09,
    D-10, D-11). Il timestamp di generazione e' scritto a coordinata fissa
    (`B4`, etichetta in `A4`) perche' il re-import (08-08) deve rileggerlo
    deterministicamente, non cercarlo in una cella qualsiasi.

    Quando `con_margine` e' True, un foglio `AVVISO` viene inserito in
    posizione 0 (prima di ISTRUZIONI): e' il primo foglio che l'admin vede
    aprendo il file (D-13, UI-SPEC §B).
    """
    wb = Workbook()
    istruzioni = wb.active
    istruzioni.title = "ISTRUZIONI"
    istruzioni["A1"] = _TESTO_REGOLA_CELLA_VUOTA
    istruzioni["A2"] = _TESTO_NO_ELIMINAZIONE
    istruzioni["A3"] = _TESTO_NO_TOCCARE_ID
    istruzioni["A4"] = "Generato il:"
    istruzioni["B4"] = datetime.now(timezone.utc).isoformat()

    if con_margine:
        avviso = wb.create_sheet("AVVISO", 0)
        avviso["A1"] = _TESTO_AVVISO_MARGINE

    return wb


def _valore_cella(valore: object) -> object:
    """`Decimal` -> `float` (vedi docstring di modulo); ogni altro tipo
    passa invariato."""
    if isinstance(valore, Decimal):
        return float(valore)
    return valore


def scrivi_foglio_dati(
    wb: Workbook, *, nome: str, colonne: list[str], righe: list[dict]
) -> int:
    """Crea un foglio `nome` con header esatto `colonne` (prima colonna
    sempre `id`, D-03) e una riga per record. Ritorna il numero di righe
    dati scritte."""
    ws = wb.create_sheet(nome)
    ws.append(list(colonne))
    for riga in righe:
        ws.append([_valore_cella(riga.get(colonna)) for colonna in colonne])
    return len(righe)


def build_filename(risorsa_slug: str, *, con_margine: bool = False) -> str:
    """`{risorsa_slug}_{YYYY-MM-DD}.xlsx`, o `{risorsa_slug}_con-margine_
    {YYYY-MM-DD}.xlsx` nella variante marcata (UI-SPEC §B)."""
    data_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    marcatore = "_con-margine" if con_margine else ""
    return f"{risorsa_slug}{marcatore}_{data_str}.xlsx"


def serializza(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
