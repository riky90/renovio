"""Estrazione dei computi reali (`docs/casi-reali/*.xlsx`) in fixture JSON per
la taratura del computo parametrico (TAR-01, fase 11).

Per ogni caso scrive `backend/tests/fixtures/casi_reali/caso-0N.json` con:

- `totali`: i totali del foglio `Analisi preliminareTOTALI` (IVA inclusa) ed
  extra (CSE/CSP, pratiche);
- `voci`: ogni riga con quantita' > 0 dei fogli A/B/C/D, con codice del seed,
  UM, quantita', prezzo cliente e totale cliente.

Le righe di A-EDILE prendono il codice dalla stessa macchina a stati del seed
(`computo_parsing.extract_a_edile`), poi il codice e' verificato sulla
descrizione del seed: un file vecchio puo' avere codici diversi (lo
smaltimento `A.9.5` del 2025 e' l'`A.9.6` di oggi). Una riga che il tecnico ha
riscritto (stesso posto, descrizione diversa) tiene il codice del posto ed e'
marcata `descrizione_modificata`. Le righe di B/C/D si riconoscono per
descrizione (e prezzo, se la descrizione non basta).

MAI le colonne impresa (TOTALE IMPRESA, coefficiente, fee, guadagno): regola
di dominio 1. Lo script legge solo le colonne cliente (UM, Qta, Prezzo U.,
Totale) e non scrive le note del tecnico.

Gli `input` del wizard (m², tipo, risposte) non si estraggono: sono una
deduzione, scritta a mano nella fixture e preservata dallo script quando la
fixture esiste gia'.

Run (da backend/): `.venv/bin/python scripts/estrai_casi_reali.py`
"""

from __future__ import annotations

import json
import re
import sys
import warnings
from pathlib import Path
from unittest import mock

_BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))

import openpyxl  # noqa: E402

from app.services import computo_parsing  # noqa: E402

CASI_DIR = _BACKEND_DIR.parent / "docs" / "casi-reali"
SEED_PATH = _BACKEND_DIR / "seed_data" / "seed-dettaglio-pro.json"
OUT_DIR = _BACKEND_DIR / "tests" / "fixtures" / "casi_reali"

FOGLI_CATALOGO = {
    "B-FINITURE": "B-FINITURE",
    "C- SERRAMENTI PORTE  ": "C-SERRAMENTI",
    "D- PRODOTTI": "D-PRODOTTI",
}
#: Righe del foglio TOTALI: etichetta (inizio) -> chiave della fixture.
TOTALI = {
    "A-EDILE": "A-EDILE",
    "B-FINITURE": "B-FINITURE",
    "C-PORTE E SERRAMENTI": "C-SERRAMENTI",
    "D-PRODOTTI": "D-PRODOTTI",
    "NOMINA DI CSE/CSP": "cse_csp",
    "PRATICHE AMMINISTRATIVE": "pratiche",
    "TOTALE  iva inclusa": "totale_lavori",
    "da pagare iva INCLUSA": "da_pagare",
}


def _norm(testo: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(testo).lower())


def _num(valore) -> float | None:
    if isinstance(valore, bool) or not isinstance(valore, (int, float)):
        return None
    return round(float(valore), 4)


def _seed_voci() -> dict[str, list[dict]]:
    dati = json.loads(SEED_PATH.read_text(encoding="utf-8"))
    per_cat: dict[str, list[dict]] = {}
    for s in dati["sottocategorie"]:
        for v in s["voci"]:
            per_cat.setdefault(s["categoria_code"], []).append(v)
    return per_cat


def _simile(a: str, b: str, n: int = 60) -> bool:
    return _norm(a)[:n] == _norm(b)[:n]


def _righe_a_edile(ws) -> list[dict]:
    """(codice del seed, riga Excel, descrizione) di ogni voce di A-EDILE,
    con la stessa macchina a stati del seed."""
    catturate: list[dict] = []
    originale = computo_parsing.append_voce

    def cattura(sotto, codice, descrizione, um, prezzo, ordine, scarti=None, riga_excel=None):
        voce = originale(sotto, codice, descrizione, um, prezzo, ordine, scarti, riga_excel)
        if voce is not None:
            catturate.append({"codice": codice, "riga": riga_excel, "descrizione": descrizione})
        return voce

    with mock.patch.object(computo_parsing, "append_voce", cattura):
        computo_parsing.extract_a_edile(ws, scarti=[])
    return catturate


def _estrai_a_edile(ws, seed_a: list[dict]) -> list[dict]:
    per_codice = {v["codice"]: v for v in seed_a}
    voci = []
    for r in _righe_a_edile(ws):
        riga = [c.value for c in ws[r["riga"]]]
        quantita, totale = _num(riga[4]), _num(riga[6])
        if not quantita and not totale:
            continue
        codice = r["codice"]
        if codice and codice.endswith("."):
            codice = codice.rstrip(".")  # "A.12.1." nei file vecchi
        voce = {
            "foglio": "A-EDILE",
            "codice": codice,
            "riga": r["riga"],
            "impianto": riga[0],
            "descrizione": r["descrizione"][:140],
            "um": str(riga[3]).strip(),
            "quantita": quantita,
            "prezzo": _num(riga[5]),
            "totale": totale,
        }
        if quantita is None:
            # Quantita' scritta a parole ("T" per piano terra): conta il totale.
            voce["quantita_testo"] = str(riga[4])
        seed = per_codice.get(codice)
        if seed is None or not _simile(seed["descrizione"], r["descrizione"]):
            candidati = [v for v in seed_a if _simile(v["descrizione"], r["descrizione"])]
            if len(candidati) == 1:
                voce["codice_nel_file"] = codice
                voce["codice"] = candidati[0]["codice"]
            else:
                voce["descrizione_modificata"] = True
        voci.append(voce)
    return voci


def _estrai_catalogo(ws, categoria: str, seed_cat: list[dict]) -> list[dict]:
    voci = []
    for idx, riga in enumerate(ws.iter_rows(values_only=True), start=1):
        descrizione, um, quantita, prezzo = riga[0], riga[1], _num(riga[2]), _num(riga[3])
        if descrizione is None or um is None or not quantita:
            continue
        descrizione = str(descrizione).strip().replace("\n", " ")
        candidati = [v for v in seed_cat if _simile(v["descrizione"], descrizione, 200)]
        if len(candidati) > 1:
            candidati = [v for v in candidati if abs(v["prezzo_unitario_cliente"] - (prezzo or 0)) < 0.01] or candidati
        voce = {
            "foglio": categoria,
            "codice": candidati[0]["codice"] if len(candidati) == 1 else None,
            "riga": idx,
            "descrizione": descrizione[:140],
            "um": str(um).strip(),
            "quantita": quantita,
            "prezzo": prezzo,
            "totale": _num(riga[4]),
        }
        if len(candidati) != 1:
            voce["codici_candidati"] = [v["codice"] for v in candidati]
        voci.append(voce)
    return voci


def _totali(ws) -> dict:
    out: dict[str, float] = {}
    for riga in ws.iter_rows(values_only=True):
        etichetta = next((str(c).strip() for c in riga[:2] if isinstance(c, str)), None)
        valore = _num(riga[5]) if len(riga) > 5 else None
        if etichetta is None or valore is None:
            continue
        for prefisso, chiave in TOTALI.items():
            if etichetta.startswith(prefisso):
                out[chiave] = round(valore, 2)
    return out


def _descrizione_sintetica(ws) -> str | None:
    """La prima cella piena sotto "DESCRIZIONE SINTETICA DELLE OPERE"."""
    trovata = False
    for riga in ws.iter_rows(min_row=1, max_row=30, values_only=True):
        testo = riga[1] if isinstance(riga[1], str) else None
        if testo and testo.startswith("DESCRIZIONE SINTETICA"):
            trovata = True
        elif trovata and testo:
            return " ".join(testo.split())
    return None


def estrai(path: Path, seed: dict[str, list[dict]]) -> dict:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    voci = _estrai_a_edile(wb["A-EDILE"], seed["A-EDILE"])
    for foglio, categoria in FOGLI_CATALOGO.items():
        voci += _estrai_catalogo(wb[foglio], categoria, seed.get(categoria, []))

    somme: dict[str, float] = {}
    for v in voci:
        somme[v["foglio"]] = round(somme.get(v["foglio"], 0) + (v["totale"] or 0), 2)
    return {
        "file": f"docs/casi-reali/{path.name}",
        "descrizione": _descrizione_sintetica(wb["A-EDILE"]),
        "totali_ivati": _totali(wb["Analisi preliminareTOTALI"]),
        "somme_imponibili_dalle_voci": somme,
        "voci": voci,
    }


def main() -> None:
    seed = _seed_voci()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for path in sorted(CASI_DIR.glob("caso-*.xlsx")):
        caso = path.name[: len("caso-0N")]
        out = OUT_DIR / f"{caso}.json"
        dati = {"caso": caso, **estrai(path, seed)}
        if out.exists():
            precedente = json.loads(out.read_text(encoding="utf-8"))
            for chiave in ("input", "deduzioni", "soglia"):
                if chiave in precedente:
                    dati[chiave] = precedente[chiave]
        out.write_text(json.dumps(dati, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        senza = [v for v in dati["voci"] if v["codice"] is None or v.get("descrizione_modificata")]
        print(f"{caso}: {len(dati['voci'])} voci, {len(senza)} da controllare -> {out}")


if __name__ == "__main__":
    main()
