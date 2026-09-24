"""Test che contano per `/admin/import-excel/domande` (08.1-05-PLAN.md,
CUR-01/CUR-02): round-trip a vuoto sui cinque fogli scritti dal piano 03
(`domande`, `opzioni`, `effetti`, `attese`, `derivazioni`), la guardia D-28
attraversata dal file (non aggirabile ricaricando un export), `attiva`/
`attivo` che si cambiano dal file su tutti e cinque i fogli (D-30), la
derivazione nuova creata dal foglio senza toccare il seed (D-06), la
domanda nuova con la sua opzione implicita per gli stepper (D-35), il
gate admin-only, e la marcatura di staleness (D-11).
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine

import app.db as db_module
from app.config import settings
from app.db import Base
from app.limiter import limiter
from app.main import app
from app.models.domanda import (
    AttesaImmobile,
    DerivazioneRisposta,
    Domanda,
    EffettoOpzione,
    OpzioneDomanda,
)
from app.models.user import User
from app.services.export_excel.domande import build_workbook_domande
from app.services.export_excel.workbook import TOKEN_CANCELLA
from scripts.seed import load_seed_data, seed
from tests.session_factory import sessione_di_test

REGISTRAZIONE_URL = "/api/v1/auth/registrati"
IMPORT_URL = "/api/v1/admin/import-excel/domande"

ADMIN_EMAIL = "admin.import.domande@example.com"

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def client(monkeypatch, tmp_path):
    limiter.reset()
    monkeypatch.setattr(settings, "ADMIN_EMAIL", ADMIN_EMAIL)

    db_path = tmp_path / "test_import_domande.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    testing_session_local = sessione_di_test(engine)

    with testing_session_local() as s:
        seed(s, load_seed_data())

    monkeypatch.setattr(db_module, "SessionLocal", testing_session_local)

    import app.dependencies.auth as deps_auth_module
    import app.routers.auth as auth_router_module

    monkeypatch.setattr(deps_auth_module, "SessionLocal", testing_session_local)
    monkeypatch.setattr(auth_router_module, "SessionLocal", testing_session_local)

    test_client = TestClient(app)
    test_client._session_local = testing_session_local
    return test_client


def _registra_admin(client, email=ADMIN_EMAIL, nome="Admin Import Domande"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _registra_ruolo(client, role: str, email: str, nome: str):
    resp = client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})
    with client._session_local() as db:
        user = db.query(User).filter_by(email=email).one()
        user.role = role
        db.commit()
    return resp


def _workbook_default(client) -> bytes:
    with client._session_local() as db:
        file_bytes, _ = build_workbook_domande(db)
    return file_bytes


def _post_preview(client, xlsx_bytes: bytes):
    return client.post(
        f"{IMPORT_URL}/preview",
        files={"file": ("domande.xlsx", xlsx_bytes, MIME_XLSX)},
    )


def _a_riga_piatta_in(diff_riga: dict) -> dict:
    """Ricostruisce `RigaPiattaIn` da `DiffRigaOut` esattamente come fa il
    client (`FlatImportDialog.tsx::toRigaPiattaIn`, CR-02): un `None` in
    `valori_nuovi` e' una cancellazione e viaggia in `campi_da_cancellare`,
    mai dentro `valori`."""
    valori = {}
    campi_da_cancellare = []
    for campo, valore in diff_riga["valori_nuovi"].items():
        if valore is None:
            campi_da_cancellare.append(campo)
        else:
            valori[campo] = valore
    return {
        "foglio": diff_riga["foglio"],
        "id": diff_riga["id"],
        "valori": valori,
        "campi_da_cancellare": campi_da_cancellare,
    }


def _post_confirm(client, preview_body: dict):
    nuove = [_a_riga_piatta_in(r) for r in preview_body["nuove"]]
    modificate = [_a_riga_piatta_in(r) for r in preview_body["modificate"]]
    return client.post(
        f"{IMPORT_URL}/confirm",
        json={"risorsa": "domande", "nuove": nuove, "modificate": modificate},
    )


def _modifica_cella(xlsx_bytes: bytes, foglio: str, riga_id, colonna: str, nuovo_valore) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[foglio]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_id = header.index("id")
    col_target = header.index(colonna)
    for row in ws.iter_rows(min_row=2):
        if str(row[col_id].value) == str(riga_id):
            row[col_target].value = nuovo_valore
            break
    else:
        raise AssertionError(f"riga con id={riga_id!r} non trovata nel foglio {foglio}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rimuovi_colonna(xlsx_bytes: bytes, foglio: str, colonna: str) -> bytes:
    """Cancella la colonna dall'header (CUR-02): `leggi_foglio_piatto`
    ignora una colonna il cui nome-header e' assente, esattamente come se
    l'admin l'avesse tolta prima di ricaricare il file."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[foglio]
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value == colonna:
            cell.value = None
            break
    else:
        raise AssertionError(f"colonna {colonna!r} non trovata nel foglio {foglio}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _tronca_righe(xlsx_bytes: bytes, foglio: str, mantieni: int) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[foglio]
    for r in range(ws.max_row, 1 + mantieni, -1):
        ws.delete_rows(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _aggiungi_riga(xlsx_bytes: bytes, foglio: str, valori: dict) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[foglio]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    riga = ws.max_row + 1
    for colonna, valore in valori.items():
        ws.cell(row=riga, column=header.index(colonna) + 1, value=valore)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestRoundtripAVuoto:
    """D-12: rileggere l'export appena scaricato non deve proporre nulla —
    ne' una creazione, ne' una modifica, ne' uno scarto — e il numero di
    righe lette per foglio deve coincidere col conteggio della tabella."""

    def test_roundtrip_conta_le_cinque_tabelle_e_la_preview_e_vuota(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        from app.services.import_excel.registry import risolvi_handler

        handler = risolvi_handler("domande")
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        righe, scarti_parse = handler.parse(wb)
        assert scarti_parse == []

        per_foglio: dict[str, int] = {}
        for r in righe:
            per_foglio[r.foglio] = per_foglio.get(r.foglio, 0) + 1

        with client._session_local() as db:
            conteggi = {
                "domande": db.query(Domanda).count(),
                "opzioni": db.query(OpzioneDomanda).count(),
                "effetti": db.query(EffettoOpzione).count(),
                "attese": db.query(AttesaImmobile).count(),
                "derivazioni": db.query(DerivazioneRisposta).count(),
            }

        assert sum(conteggi.values()) > 0
        assert per_foglio.get("domande", 0) == conteggi["domande"]
        assert per_foglio.get("opzioni", 0) == conteggi["opzioni"]
        assert per_foglio.get("effetti", 0) == conteggi["effetti"]
        assert per_foglio.get("attese", 0) == conteggi["attese"]
        assert per_foglio.get("derivazioni", 0) == conteggi["derivazioni"]

        preview = _post_preview(client, xlsx)
        assert preview.status_code == 200
        body = preview.json()
        assert body["nuove"] == []
        assert body["modificate"] == []
        assert body["scarti"] == []


class TestModificaTestoDomanda:
    def test_modifica_testo_domanda_si_applica(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            code = db.query(Domanda).order_by(Domanda.code).first().code

        xlsx_mod = _modifica_cella(xlsx, "domande", code, "testo", "Testo cambiato dal file")

        preview = _post_preview(client, xlsx_mod).json()
        modificate = [r for r in preview["modificate"] if r["foglio"] == "domande" and r["id"] == code]
        assert len(modificate) == 1
        assert modificate[0]["campi_modificati"] == ["testo"]

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            assert db.get(Domanda, code).testo == "Testo cambiato dal file"


class TestColonnaAssenteLasciaIntattoIlCampo:
    """CUR-02: una colonna tolta dal file non deve toccare quel campo a
    database, anche quando la stessa riga porta un altro campo cambiato."""

    def test_colonna_etichetta_assente_lascia_intatta_letichetta(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            opzione = db.query(OpzioneDomanda).order_by(OpzioneDomanda.id).first()
            opzione_id = opzione.id
            etichetta_prima = opzione.etichetta
            ordine_prima = opzione.ordine

        xlsx_senza_etichetta = _rimuovi_colonna(xlsx, "opzioni", "etichetta")
        nuovo_ordine = ordine_prima + 1
        xlsx_mod = _modifica_cella(xlsx_senza_etichetta, "opzioni", opzione_id, "ordine", nuovo_ordine)

        preview = _post_preview(client, xlsx_mod).json()
        modificate = [
            r for r in preview["modificate"] if r["foglio"] == "opzioni" and r["id"] == str(opzione_id)
        ]
        assert len(modificate) == 1
        assert modificate[0]["campi_modificati"] == ["ordine"]
        assert "etichetta" not in modificate[0]["valori_nuovi"]

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            dopo = db.get(OpzioneDomanda, opzione_id)
            assert dopo.etichetta == etichetta_prima
            assert dopo.ordine == nuovo_ordine


class TestSottoinsiemeNonEliminaNulla:
    """D-10: un file con un sottoinsieme delle righe applicato lascia
    invariati i conteggi di tutte e cinque le tabelle."""

    def test_file_con_meta_righe_per_foglio_non_elimina_nulla(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            conteggi_prima = {
                "domande": db.query(Domanda).count(),
                "opzioni": db.query(OpzioneDomanda).count(),
                "effetti": db.query(EffettoOpzione).count(),
                "attese": db.query(AttesaImmobile).count(),
                "derivazioni": db.query(DerivazioneRisposta).count(),
            }

        xlsx_troncato = xlsx
        for foglio, totale in conteggi_prima.items():
            mantieni = max(1, totale // 2)
            xlsx_troncato = _tronca_righe(xlsx_troncato, foglio, mantieni=mantieni)

        preview = _post_preview(client, xlsx_troncato).json()
        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            conteggi_dopo = {
                "domande": db.query(Domanda).count(),
                "opzioni": db.query(OpzioneDomanda).count(),
                "effetti": db.query(EffettoOpzione).count(),
                "attese": db.query(AttesaImmobile).count(),
                "derivazioni": db.query(DerivazioneRisposta).count(),
            }

        assert conteggi_dopo == conteggi_prima


class TestD28ViaImport:
    """D-28/D-30: disattivare da file un'opzione in uso (`non_so`, insieme
    default della sua domanda e trigger di tre derivazioni) e' uno scarto
    col nome della dipendenza, non un'eccezione — e le altre righe del file
    si importano comunque."""

    def test_disattivare_unopzione_in_uso_e_uno_scarto_le_altre_righe_passano(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            non_so = (
                db.query(OpzioneDomanda)
                .filter_by(domanda_code="classe_energetica_attuale", code="non_so")
                .one()
            )
            non_so_id = non_so.id
            altra = (
                db.query(OpzioneDomanda)
                .filter(OpzioneDomanda.id != non_so_id)
                .order_by(OpzioneDomanda.id)
                .first()
            )
            altra_id = altra.id
            altra_ordine_prima = altra.ordine

        xlsx_mod = _modifica_cella(xlsx, "opzioni", non_so_id, "attiva", False)
        xlsx_mod = _modifica_cella(xlsx_mod, "opzioni", altra_id, "ordine", altra_ordine_prima + 1)

        preview = _post_preview(client, xlsx_mod)
        assert preview.status_code == 200
        body = preview.json()

        scarti_derivazione = [s for s in body["scarti"] if "derivazione" in s["motivo"]]
        assert len(scarti_derivazione) == 1

        modificate_altra = [r for r in body["modificate"] if r["id"] == str(altra_id)]
        assert len(modificate_altra) == 1

        confirm = _post_confirm(client, body)
        assert confirm.status_code == 200

        with client._session_local() as db:
            assert db.get(OpzioneDomanda, non_so_id).attiva is True
            assert db.get(OpzioneDomanda, altra_id).ordine == altra_ordine_prima + 1


class TestAttivaSiCambianoDaFile:
    """D-30: `attiva`/`attivo` si cambiano dal file su tutti e cinque i
    fogli, quando la disattivazione non urta la guardia D-28."""

    def test_attiva_attivo_si_cambiano_su_domande_effetti_attese_derivazioni(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            domanda = db.query(Domanda).filter_by(code="superficie_calpestabile").one()
            effetto = db.query(EffettoOpzione).order_by(EffettoOpzione.id).first()
            attesa = db.query(AttesaImmobile).order_by(AttesaImmobile.id).first()
            derivazione = db.query(DerivazioneRisposta).order_by(DerivazioneRisposta.id).first()
            domanda_code = domanda.code
            effetto_id = effetto.id
            attesa_id = attesa.id
            derivazione_id = derivazione.id

        xlsx_mod = _modifica_cella(xlsx, "domande", domanda_code, "attiva", False)
        xlsx_mod = _modifica_cella(xlsx_mod, "effetti", effetto_id, "attivo", False)
        xlsx_mod = _modifica_cella(xlsx_mod, "attese", attesa_id, "attiva", False)
        xlsx_mod = _modifica_cella(xlsx_mod, "derivazioni", derivazione_id, "attiva", False)

        preview = _post_preview(client, xlsx_mod).json()
        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            assert db.get(Domanda, domanda_code).attiva is False
            assert db.get(EffettoOpzione, effetto_id).attivo is False
            assert db.get(AttesaImmobile, attesa_id).attiva is False
            assert db.get(DerivazioneRisposta, derivazione_id).attiva is False


class TestDerivazioneNuovaDaFoglio:
    """D-06: una derivazione nuova si crea dal foglio, senza toccare il
    seed; un codice inesistente e' uno scarto descritto."""

    def test_riga_senza_id_con_codici_validi_crea_derivazione(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        xlsx_mod = _aggiungi_riga(
            xlsx,
            "derivazioni",
            {
                "domanda_code": "classe_energetica_attuale",
                "quando_opzione_code": "a",
                "da_domanda_code": "anno_costruzione",
                "da_opzione_code": "pre_1980",
                "opzione_code": "b",
            },
        )

        preview = _post_preview(client, xlsx_mod).json()
        nuove = [r for r in preview["nuove"] if r["foglio"] == "derivazioni"]
        assert len(nuove) == 1

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            creata = (
                db.query(DerivazioneRisposta)
                .filter_by(
                    domanda_code="classe_energetica_attuale",
                    quando_opzione_code="a",
                    da_domanda_code="anno_costruzione",
                    da_opzione_code="pre_1980",
                )
                .one()
            )
            assert creata.opzione_code == "b"

    def test_riga_con_da_opzione_code_inesistente_e_uno_scarto_descritto(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        xlsx_mod = _aggiungi_riga(
            xlsx,
            "derivazioni",
            {
                "domanda_code": "classe_energetica_attuale",
                "quando_opzione_code": "b",
                "da_domanda_code": "anno_costruzione",
                "da_opzione_code": "codice_fantasma",
                "opzione_code": "c",
            },
        )

        preview = _post_preview(client, xlsx_mod).json()
        assert [r for r in preview["nuove"] if r["foglio"] == "derivazioni"] == []
        scarti = [s for s in preview["scarti"] if "derivazioni" in s["motivo"]]
        assert len(scarti) == 1

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200
        with client._session_local() as db:
            assert (
                db.query(DerivazioneRisposta)
                .filter_by(da_opzione_code="codice_fantasma")
                .count()
                == 0
            )


class TestDomandaNuovaDaFoglio:
    """D-35: un `id` (= code) mai visto crea una domanda nuova; se
    `tipo_controllo` e' `stepper`, l'opzione implicita `valore` esiste dopo
    il confirm."""

    def test_riga_con_id_mai_visto_crea_domanda_nuova_con_opzione_valore(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        xlsx_mod = _aggiungi_riga(
            xlsx,
            "domande",
            {
                "id": "domanda_di_prova",
                "testo": "Domanda di prova",
                "tipo_controllo": "stepper",
                "gruppo": "prova",
                "ordine": 999,
                "tier_minimo": "free",
            },
        )

        preview = _post_preview(client, xlsx_mod).json()
        nuove = [r for r in preview["nuove"] if r["foglio"] == "domande" and r["id"] == "domanda_di_prova"]
        assert len(nuove) == 1

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            creata = db.get(Domanda, "domanda_di_prova")
            assert creata is not None
            assert creata.tipo_controllo == "stepper"
            opzione_valore = (
                db.query(OpzioneDomanda)
                .filter_by(domanda_code="domanda_di_prova", code="valore")
                .one_or_none()
            )
            assert opzione_valore is not None


class TestEffettoNuovoDaFoglio:
    def test_riga_senza_id_con_domanda_e_opzione_valida_crea_effetto(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            opzione = (
                db.query(OpzioneDomanda)
                .filter_by(domanda_code="anno_costruzione", code="post_1990")
                .one()
            )
            conteggio_prima = db.query(EffettoOpzione).count()

        xlsx_mod = _aggiungi_riga(
            xlsx,
            "effetti",
            {
                "domanda_code": "anno_costruzione",
                "opzione_code": "post_1990",
                "tipo_effetto": "moltiplicatore_categoria",
                "target_categoria_code": "B-FINITURE",
                "valore": 1.05,
            },
        )

        preview = _post_preview(client, xlsx_mod).json()
        nuove = [r for r in preview["nuove"] if r["foglio"] == "effetti"]
        assert len(nuove) == 1

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            assert db.query(EffettoOpzione).count() == conteggio_prima + 1
            creato = (
                db.query(EffettoOpzione)
                .filter_by(opzione_id=opzione.id, target_categoria_code="B-FINITURE")
                .one()
            )
            assert float(creato.valore) == pytest.approx(1.05)

    def test_riga_con_bersaglio_non_ammesso_e_uno_scarto_col_messaggio_di_valida_effetto(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            conteggio_prima = db.query(EffettoOpzione).count()

        xlsx_mod = _aggiungi_riga(
            xlsx,
            "effetti",
            {
                "domanda_code": "anno_costruzione",
                "opzione_code": "post_1990",
                "tipo_effetto": "moltiplicatore_categoria",
                "target_voce_id": 1,
                "valore": 1.05,
            },
        )

        preview = _post_preview(client, xlsx_mod).json()
        assert [r for r in preview["nuove"] if r["foglio"] == "effetti"] == []
        scarti = [s for s in preview["scarti"] if "non puo' puntare" in s["motivo"]]
        assert len(scarti) == 1

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200
        with client._session_local() as db:
            assert db.query(EffettoOpzione).count() == conteggio_prima


class TestEffettoIdentitaNonSiCambia:
    def test_cambiare_tipo_effetto_e_uno_scarto_non_un_aggiornamento(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            effetto = (
                db.query(EffettoOpzione)
                .filter_by(tipo_effetto="moltiplicatore_categoria")
                .order_by(EffettoOpzione.id)
                .first()
            )
            effetto_id = effetto.id
            tipo_prima = effetto.tipo_effetto

        xlsx_mod = _modifica_cella(xlsx, "effetti", effetto_id, "tipo_effetto", "eur_mq_categoria")

        preview = _post_preview(client, xlsx_mod).json()
        assert [r for r in preview["modificate"] if r["id"] == str(effetto_id)] == []
        scarti = [s for s in preview["scarti"] if "tipo e bersaglio non si cambiano" in s["motivo"]]
        assert len(scarti) == 1

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            assert db.get(EffettoOpzione, effetto_id).tipo_effetto == tipo_prima


class TestCancellaConToken:
    def test_cancella_su_help_lo_azzera(self, client):
        _registra_admin(client)
        with client._session_local() as db:
            domanda = db.query(Domanda).filter(Domanda.help.isnot(None)).first()
            assert domanda is not None, "seed atteso: almeno una domanda con help"
            code = domanda.code

        xlsx = _workbook_default(client)
        xlsx_mod = _modifica_cella(xlsx, "domande", code, "help", TOKEN_CANCELLA)

        preview = _post_preview(client, xlsx_mod).json()
        modificate = [r for r in preview["modificate"] if r["id"] == code]
        assert len(modificate) == 1
        assert modificate[0]["valori_nuovi"]["help"] is None

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            assert db.get(Domanda, code).help is None

    def test_cancella_su_testo_e_uno_scarto(self, client):
        _registra_admin(client)
        xlsx = _workbook_default(client)

        with client._session_local() as db:
            code = db.query(Domanda).order_by(Domanda.code).first().code
            testo_prima = db.get(Domanda, code).testo

        xlsx_mod = _modifica_cella(xlsx, "domande", code, "testo", TOKEN_CANCELLA)

        preview = _post_preview(client, xlsx_mod).json()
        assert [r for r in preview["modificate"] if r["id"] == code] == []
        scarti = [s for s in preview["scarti"] if "non annullabile" in s["motivo"]]
        assert len(scarti) == 1

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200
        with client._session_local() as db:
            assert db.get(Domanda, code).testo == testo_prima


class TestRigaModificataDopoExportEMarcata:
    def test_riga_modificata_dopo_export_e_marcata(self, client):
        """D-11: due righe cambiate a database dopo l'export, una con
        `updated_at` successivo al file e una con `updated_at` precedente —
        entrambe compaiono fra le modificate, solo la prima porta
        `riga_modificata_dopo_export`."""
        _registra_admin(client)
        xlsx = _workbook_default(client)  # generato_il ~= ora

        with client._session_local() as db:
            righe = db.query(Domanda).order_by(Domanda.code).all()
            assert len(righe) >= 2

            riga_stale, riga_non_stale = righe[0], righe[1]
            riga_stale.gruppo = f"{riga_stale.gruppo}-mod"
            riga_stale.updated_at = datetime.now(timezone.utc) + timedelta(hours=1)

            riga_non_stale.gruppo = f"{riga_non_stale.gruppo}-mod"
            riga_non_stale.updated_at = datetime(2020, 1, 1, tzinfo=timezone.utc)

            db.commit()

        preview = _post_preview(client, xlsx).json()
        modificate = [r for r in preview["modificate"] if r["foglio"] == "domande"]
        assert len(modificate) == 2

        marcate = [r for r in modificate if r["riga_modificata_dopo_export"]]
        assert len(marcate) == 1


class TestImportVietatoAlNonAdmin:
    def test_import_vietato_al_non_admin(self, client):
        _registra_admin(client, email="admin.per.file.domande@example.com", nome="Admin Per File")
        xlsx = _workbook_default(client)

        for role, email in (
            ("consulente", "consulente.import.domande@example.com"),
            ("user", "utente.import.domande@example.com"),
        ):
            _registra_ruolo(client, role, email, f"Ruolo {role}")

            response_preview = _post_preview(client, xlsx)
            assert response_preview.status_code == 403

            response_confirm = client.post(
                f"{IMPORT_URL}/confirm", json={"risorsa": "domande", "nuove": [], "modificate": []}
            )
            assert response_confirm.status_code == 403
