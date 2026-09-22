"""E2E tests (08-01-PLAN.md, T-08-01..T-08-04) per
`GET /api/v1/admin/listini/{listino_id}/export.xlsx`.

Copre: 200 per admin con MIME e Content-Disposition corretti; il workbook
riletto con openpyxl ha il foglio `ISTRUZIONI` con `#CANCELLA#` e un
timestamp parsabile, e il foglio `voci` con l'header esatto delle dodici
colonne; il numero di righe dati coincide col conteggio DB delle voci del
listino; 403 per un utente non admin e nessuna riga di log scritta; una
riga `export_log` dopo il download riuscito.
"""

from __future__ import annotations

import io
from datetime import datetime

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import app.db as db_module
from app.config import settings
from app.db import Base
from app.limiter import limiter
from app.main import app
from app.models.dettaglio_pro import Sottocategoria, Voce
from app.models.export_log import ExportLog
from scripts.seed import load_seed_data, seed

REGISTRAZIONE_URL = "/api/v1/auth/registrati"
LISTINI_URL = "/api/v1/admin/listini"

ADMIN_EMAIL = "admin.export@example.com"


@pytest.fixture()
def client(monkeypatch, tmp_path):
    limiter.reset()
    monkeypatch.setattr(settings, "ADMIN_EMAIL", ADMIN_EMAIL)

    db_path = tmp_path / "test_export_excel_listino.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    testing_session_local = sessionmaker(bind=engine)

    with testing_session_local() as s:
        seed(s, load_seed_data())

    monkeypatch.setattr(db_module, "SessionLocal", testing_session_local)

    import app.dependencies.auth as deps_auth_module
    import app.routers.auth as auth_router_module

    monkeypatch.setattr(deps_auth_module, "SessionLocal", testing_session_local)
    monkeypatch.setattr(auth_router_module, "SessionLocal", testing_session_local)

    test_client = TestClient(app)
    test_client._session_local = testing_session_local  # per le assert dirette a DB
    return test_client


def _registra_admin(client, email=ADMIN_EMAIL, nome="Admin Export"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _registra_user(client, email="utente.export@example.com", nome="Utente Normale"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _re_listino_id(client) -> int:
    rows = client.get(LISTINI_URL).json()
    return next(r["id"] for r in rows if r["citta_riferimento"] == "Reggio Emilia")


HEADER_ATTESO = [
    "id",
    "categoria_code",
    "sottocategoria_codice",
    "sottocategoria_nome",
    "sottocategoria_famiglia",
    "sottocategoria_ordine",
    "codice",
    "descrizione",
    "um",
    "prezzo_unitario_cliente",
    "ordine",
    "livello",
]


class TestExportListino:
    def test_export_200_mime_e_content_disposition(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)

        response = client.get(f"{LISTINI_URL}/{listino_id}/export.xlsx")

        assert response.status_code == 200
        assert response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert 'attachment; filename="listino_' in response.headers["content-disposition"]

    def test_foglio_istruzioni_e_header_voci(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)

        response = client.get(f"{LISTINI_URL}/{listino_id}/export.xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(response.content))

        assert wb.sheetnames[0] == "ISTRUZIONI"
        istruzioni = wb["ISTRUZIONI"]
        testo_foglio = " ".join(
            str(cell.value) for row in istruzioni.iter_rows() for cell in row if cell.value is not None
        )
        assert "#CANCELLA#" in testo_foglio

        timestamp_cell = istruzioni["B4"].value
        datetime.fromisoformat(timestamp_cell)  # non solleva -> parsabile

        assert "voci" in wb.sheetnames
        voci_sheet = wb["voci"]
        header = [cell.value for cell in next(voci_sheet.iter_rows(min_row=1, max_row=1))]
        assert header == HEADER_ATTESO

    def test_numero_righe_coincide_con_db(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)

        response = client.get(f"{LISTINI_URL}/{listino_id}/export.xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        voci_sheet = wb["voci"]
        n_righe_file = voci_sheet.max_row - 1  # meno l'header

        with client._session_local() as db:
            n_righe_db = (
                db.query(Voce)
                .join(Sottocategoria, Voce.sottocategoria_id == Sottocategoria.id)
                .filter(Sottocategoria.listino_id == listino_id)
                .count()
            )

        assert n_righe_file == n_righe_db
        assert n_righe_file > 0

    def test_403_per_non_admin_nessuna_riga_di_log(self, client):
        _registra_user(client)
        # un utente non admin non passa nemmeno da `require_admin`: l'id
        # nel path e' irrilevante, il gate scatta prima di toccare il DB.
        response = client.get(f"{LISTINI_URL}/1/export.xlsx")

        assert response.status_code == 403
        with client._session_local() as db:
            assert db.query(ExportLog).count() == 0

    def test_download_admin_scrive_riga_export_log(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)

        response = client.get(f"{LISTINI_URL}/{listino_id}/export.xlsx")
        assert response.status_code == 200

        with client._session_local() as db:
            righe_db = (
                db.query(Voce)
                .join(Sottocategoria, Voce.sottocategoria_id == Sottocategoria.id)
                .filter(Sottocategoria.listino_id == listino_id)
                .count()
            )
            log = db.query(ExportLog).one()
            assert log.risorsa == "listino"
            assert log.include_margine is False
            assert log.righe == righe_db
