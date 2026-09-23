"""Test che contano per `/admin/import-excel/lead` (08-07-PLAN.md,
CUR-01/CUR-02/CUR-04): round-trip a vuoto, `stato`/`nota` aggiornabili in
blocco, i dati di contatto NON riscrivibili da file, la deroga a D-04
(riga senza id = scarto, mai un lead nuovo), il permesso per risorsa
(staff sui lead, admin sui parametri) e nessuna eliminazione.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import app.db as db_module
from app.config import settings
from app.db import Base
from app.limiter import limiter
from app.main import app
from app.models.lead import Lead
from app.models.user import User
from app.services.export_excel.lead import build_workbook_lead
from app.services.export_excel.parametri import build_workbook_parametri
from app.services.export_excel.workbook import TOKEN_CANCELLA
from app.services.import_excel.lead import diff as diff_lead
from app.schemas.admin_import import RigaPiattaIn
from scripts.seed import load_seed_data, seed

REGISTRAZIONE_URL = "/api/v1/auth/registrati"
LEAD_IMPORT_URL = "/api/v1/admin/import-excel/lead"
PARAMETRI_IMPORT_URL = "/api/v1/admin/import-excel/parametri"

ADMIN_EMAIL = "admin.import.lead@example.com"

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def client(monkeypatch, tmp_path):
    limiter.reset()
    monkeypatch.setattr(settings, "ADMIN_EMAIL", ADMIN_EMAIL)

    db_path = tmp_path / "test_import_lead.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    testing_session_local = sessionmaker(bind=engine)

    monkeypatch.setattr(db_module, "SessionLocal", testing_session_local)

    import app.dependencies.auth as deps_auth_module
    import app.routers.auth as auth_router_module

    monkeypatch.setattr(deps_auth_module, "SessionLocal", testing_session_local)
    monkeypatch.setattr(auth_router_module, "SessionLocal", testing_session_local)

    test_client = TestClient(app)
    test_client._session_local = testing_session_local
    return test_client


def _registra(client, email, nome="Utente"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _registra_admin(client, email=ADMIN_EMAIL, nome="Admin Import Lead"):
    return _registra(client, email, nome)


def _registra_ruolo(client, role: str, email: str, nome: str):
    resp = _registra(client, email, nome)
    with client._session_local() as db:
        user = db.query(User).filter_by(email=email).one()
        user.role = role
        db.commit()
    return resp


def _seed_lead(session_factory, *, user_id=1, nome="Mario Rossi", email="mario.rossi@example.com") -> int:
    with session_factory() as session:
        lead = Lead(
            user_id=user_id,
            nome=nome,
            email=email,
            telefono="3331234567",
            preferenza_contatto="email",
            fascia_oraria="mattina",
            messaggio="Vorrei un sopralluogo.",
            citta="Reggio Emilia",
            mq=85,
            tipo="totale",
            totale_min=Decimal("30120.00"),
            totale_max=Decimal("44350.00"),
        )
        session.add(lead)
        session.commit()
        session.refresh(lead)
        return lead.id


def _workbook_lead(client) -> bytes:
    with client._session_local() as db:
        file_bytes, _ = build_workbook_lead(db)
    return file_bytes


def _post_preview(client, xlsx_bytes: bytes):
    return client.post(
        f"{LEAD_IMPORT_URL}/preview",
        files={"file": ("lead.xlsx", xlsx_bytes, MIME_XLSX)},
    )


def _a_riga_piatta_in(diff_riga: dict) -> dict:
    """Ricostruisce `RigaPiattaIn` da `DiffRigaOut` esattamente come fa il
    client (round-trip stateless, `FlatImportDialog.tsx::toRigaPiattaIn`).

    Questo helper E' il sostituto del frontend in questa suite: se diverge
    dalla funzione TypeScript, i test smettono di dire qualcosa sul flusso
    reale. Come quella, separa i `null`: un campo il cui valore nuovo e'
    `null` e' una CANCELLAZIONE e va dichiarato in `campi_da_cancellare`,
    non lasciato dentro `valori` — dove il resolver lo stringificherebbe
    (`str(None) == "None"`, CR-02).
    """
    valori = {}
    campi_da_cancellare = []
    for campo, valore in diff_riga["valori_nuovi"].items():
        if valore is None:
            campi_da_cancellare.append(campo)
        else:
            valori[campo] = valore
    return {
        "foglio": diff_riga["foglio"],
        "id": diff_riga["id"],
        "valori": valori,
        "campi_da_cancellare": campi_da_cancellare,
    }


def _post_confirm(client, preview_body: dict):
    nuove = [_a_riga_piatta_in(r) for r in preview_body["nuove"]]
    modificate = [_a_riga_piatta_in(r) for r in preview_body["modificate"]]
    return client.post(
        f"{LEAD_IMPORT_URL}/confirm",
        json={"risorsa": "lead", "nuove": nuove, "modificate": modificate},
    )


def _modifica_cella(xlsx_bytes: bytes, foglio: str, riga_id, colonna: str, nuovo_valore) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[foglio]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_id = header.index("id")
    col_target = header.index(colonna)
    for row in ws.iter_rows(min_row=2):
        if str(row[col_id].value) == str(riga_id):
            row[col_target].value = nuovo_valore
            break
    else:
        raise AssertionError(f"riga con id={riga_id!r} non trovata nel foglio {foglio}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _tronca_righe(xlsx_bytes: bytes, foglio: str, mantieni: int) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[foglio]
    for r in range(ws.max_row, 1 + mantieni, -1):
        ws.delete_rows(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestRoundtripLeadAVuoto:
    def test_roundtrip_lead_a_vuoto(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local, nome="Mario Rossi")
        _seed_lead(client._session_local, nome="Luca Bianchi", email="luca.bianchi@example.com")
        xlsx = _workbook_lead(client)

        response = _post_preview(client, xlsx)
        assert response.status_code == 200
        body = response.json()

        assert len(body["nuove"]) == 0
        assert len(body["modificate"]) == 0
        assert len(body["scarti"]) == 0


class TestStatoENotaAggiornabili:
    def test_stato_e_nota_aggiornabili(self, client):
        _registra_admin(client)
        lead_id = _seed_lead(client._session_local)
        xlsx = _workbook_lead(client)

        xlsx_modificato = _modifica_cella(xlsx, "lead", lead_id, "stato", "gestito")
        xlsx_modificato = _modifica_cella(xlsx_modificato, "lead", lead_id, "nota", "Chiamato, richiamare giovedi'")

        preview = _post_preview(client, xlsx_modificato).json()
        assert len(preview["modificate"]) == 1
        riga = preview["modificate"][0]
        assert set(riga["campi_modificati"]) == {"stato", "nota"}
        assert len(preview["scarti"]) == 0

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200
        assert confirm.json()["righe_applicate"] == 1

        with client._session_local() as db:
            riga_dopo = db.get(Lead, lead_id)
            assert riga_dopo.stato == "gestito"
            assert riga_dopo.nota == "Chiamato, richiamare giovedi'"


class TestEmailNonModificabileDaFile:
    def test_email_non_modificabile_da_file(self, client):
        _registra_admin(client)
        lead_id = _seed_lead(client._session_local, email="originale@example.com")
        xlsx = _workbook_lead(client)

        xlsx_modificato = _modifica_cella(xlsx, "lead", lead_id, "email", "manomessa@example.com")

        preview = _post_preview(client, xlsx_modificato).json()
        assert len(preview["modificate"]) == 0
        assert len(preview["scarti"]) == 1
        assert "email" in preview["scarti"][0]["motivo"]

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200
        assert confirm.json()["righe_applicate"] == 0

        with client._session_local() as db:
            riga_dopo = db.get(Lead, lead_id)
            assert riga_dopo.email == "originale@example.com"


class TestRigaSenzaIdNonCreaLead:
    def test_riga_senza_id_non_crea_lead(self, client):
        _registra_admin(client)
        lead_id = _seed_lead(client._session_local)
        xlsx = _workbook_lead(client)

        xlsx_senza_id = _modifica_cella(xlsx, "lead", lead_id, "id", None)

        with client._session_local() as db:
            conteggio_prima = db.query(Lead).count()

        preview = _post_preview(client, xlsx_senza_id).json()
        assert len(preview["nuove"]) == 0
        assert any("senza id" in s["motivo"] for s in preview["scarti"])

        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            conteggio_dopo = db.query(Lead).count()
        assert conteggio_dopo == conteggio_prima


class TestPermessoPerRisorsa:
    def test_import_lead_permesso_al_consulente(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local)
        _registra_ruolo(client, "consulente", "consulente.import.lead@example.com", "Consulente Import Lead")
        xlsx = _workbook_lead(client)

        response = _post_preview(client, xlsx)
        assert response.status_code == 200

    def test_import_lead_vietato_all_utente(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local)
        _registra_ruolo(client, "user", "utente.import.lead@example.com", "Utente Import Lead")
        xlsx = _workbook_lead(client)

        response_preview = _post_preview(client, xlsx)
        assert response_preview.status_code == 403

        response_confirm = client.post(
            f"{LEAD_IMPORT_URL}/confirm", json={"risorsa": "lead", "nuove": [], "modificate": []}
        )
        assert response_confirm.status_code == 403

    def test_consulente_non_puo_importare_parametri(self, client):
        _registra_admin(client)
        with client._session_local() as db:
            seed(db, load_seed_data())
        _registra_ruolo(client, "consulente", "consulente.no.parametri@example.com", "Consulente No Parametri")

        with client._session_local() as db:
            xlsx_parametri, _ = build_workbook_parametri(db, con_margine=False)

        response = client.post(
            f"{PARAMETRI_IMPORT_URL}/preview",
            files={"file": ("parametri.xlsx", xlsx_parametri, MIME_XLSX)},
        )
        assert response.status_code == 403


class TestImportNonEliminaMaiLead:
    def test_import_non_elimina_mai_lead(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local, nome="Mario Rossi")
        _seed_lead(client._session_local, nome="Luca Bianchi", email="luca.bianchi@example.com")
        xlsx = _workbook_lead(client)

        with client._session_local() as db:
            conteggio_prima = db.query(Lead).count()

        xlsx_troncato = _tronca_righe(xlsx, "lead", mantieni=1)

        preview = _post_preview(client, xlsx_troncato).json()
        confirm = _post_confirm(client, preview)
        assert confirm.status_code == 200

        with client._session_local() as db:
            conteggio_dopo = db.query(Lead).count()
        assert conteggio_dopo == conteggio_prima


class TestCancellazioneRoundTrip:
    """CR-02: il segnale `#CANCELLA#` deve sopravvivere al giro
    preview -> client -> confirm.

    `diff()` rappresenta una cancellazione con `valori_nuovi[campo] = None`.
    Il client ricostruisce il payload di `/confirm` da quel dict: se
    rimanda indietro il `None` dentro `valori` invece che in
    `campi_da_cancellare`, il resolver fa `str(None)` e scrive a database
    la stringa letterale `"None"` — il campo non viene cancellato, viene
    corrotto. Nessun test lo copriva perche' l'helper `_a_riga_piatta_in`
    ricostruiva il payload con lo stesso errore del frontend.
    """

    def test_cancella_nota_scrive_null_non_la_stringa_none(self, client):
        _registra_admin(client)
        lead_id = _seed_lead(client._session_local)
        with client._session_local() as db:
            db.query(Lead).filter_by(id=lead_id).update({"nota": "Da richiamare"})
            db.commit()

        xlsx = _workbook_lead(client)
        xlsx_modificato = _modifica_cella(xlsx, "lead", lead_id, "nota", TOKEN_CANCELLA)

        preview = _post_preview(client, xlsx_modificato)
        assert preview.status_code == 200
        body = preview.json()
        assert body["scarti"] == []
        assert len(body["modificate"]) == 1
        # la preview propone davvero una cancellazione
        assert body["modificate"][0]["valori_nuovi"]["nota"] is None

        confirm = _post_confirm(client, body)
        assert confirm.status_code == 200
        assert confirm.json()["righe_applicate"] == 1

        with client._session_local() as db:
            lead = db.get(Lead, lead_id)
            assert lead.nota != "None", "scritta la stringa letterale 'None' invece di NULL"
            assert lead.nota is None

    def test_un_null_dentro_valori_e_uno_scarto_non_una_stringa(self, client):
        """Difesa in profondita': anche con un client che sbaglia a
        ricostruire il payload, il backend non deve stringificare il
        `None`. La riga diventa uno scarto e il database resta com'era."""
        _registra_admin(client)
        lead_id = _seed_lead(client._session_local)
        with client._session_local() as db:
            db.query(Lead).filter_by(id=lead_id).update({"nota": "Da richiamare"})
            db.commit()

        # payload ricostruito male, come faceva il frontend prima di CR-02
        payload_sbagliato = {
            "risorsa": "lead",
            "nuove": [],
            "modificate": [
                {
                    "foglio": "lead",
                    "id": str(lead_id),
                    "valori": {"nota": None},
                    "campi_da_cancellare": [],
                }
            ],
        }
        confirm = client.post(f"{LEAD_IMPORT_URL}/confirm", json=payload_sbagliato)
        assert confirm.status_code == 200
        assert confirm.json()["righe_applicate"] == 0

        with client._session_local() as db:
            assert db.get(Lead, lead_id).nota == "Da richiamare"

    def test_la_preview_segnala_lo_scarto_su_un_null_dentro_valori(self, client):
        """Lo stesso payload malformato, visto da `diff()`: deve produrre
        uno scarto descritto, cosi' l'errore e' visibile e non silenzioso."""
        _registra_admin(client)
        lead_id = _seed_lead(client._session_local)

        with client._session_local() as db:
            righe = [
                RigaPiattaIn(foglio="lead", id=str(lead_id), valori={"nota": None}, campi_da_cancellare=[])
            ]
            _, modificate, scarti = diff_lead(db, righe)

        assert modificate == []
        assert len(scarti) == 1
        assert "nota" in scarti[0].motivo
        assert "campi_da_cancellare" in scarti[0].motivo
