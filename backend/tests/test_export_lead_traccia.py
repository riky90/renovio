"""E2E test (08-05-PLAN.md, CUR-04/D-16) per `GET /api/v1/admin/lead/export.xlsx`
e `GET /api/v1/admin/export-log`.

Copre: 200 per admin e consulente (D-14, il consulente vede gia' gli stessi
dati a schermo), 403 per un utente `role="user"`; il workbook riletto ha
l'header esatto delle diciassette colonne; il numero di righe della riga di
log coincide con `db.query(Lead).count()`; il 403 non lascia riga di log;
`Cache-Control: no-store` nella risposta; l'export-log torna in ordine
decrescente di `created_at`, limitato a 50 righe, 403 per `role="user"`, e
`user_label` non e' mai l'id nudo.
"""

from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine

import app.db as db_module
from app.config import settings
from app.db import Base
from app.limiter import limiter
from app.main import app
from app.models.export_log import ExportLog
from app.models.lead import Lead
from app.models.user import User
from tests.session_factory import sessione_di_test

REGISTRAZIONE_URL = "/api/v1/auth/registrati"
LEAD_EXPORT_URL = "/api/v1/admin/lead/export.xlsx"
EXPORT_LOG_URL = "/api/v1/admin/export-log"

ADMIN_EMAIL = "admin.export.lead@example.com"

HEADER_ATTESO = [
    "id",
    "created_at",
    "nome",
    "email",
    "telefono",
    "preferenza_contatto",
    "fascia_oraria",
    "messaggio",
    "citta",
    "mq",
    "tipo",
    "totale_min",
    "totale_max",
    "referente_id",
    "email_inviata",
    "stato",
    "nota",
]


@pytest.fixture()
def client(monkeypatch, tmp_path):
    limiter.reset()
    monkeypatch.setattr(settings, "ADMIN_EMAIL", ADMIN_EMAIL)

    db_path = tmp_path / "test_export_lead_traccia.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    testing_session_local = sessione_di_test(engine)

    monkeypatch.setattr(db_module, "SessionLocal", testing_session_local)

    import app.dependencies.auth as deps_auth_module
    import app.routers.auth as auth_router_module

    monkeypatch.setattr(deps_auth_module, "SessionLocal", testing_session_local)
    monkeypatch.setattr(auth_router_module, "SessionLocal", testing_session_local)

    for module_name in ("app.routers.admin.lead", "app.routers.admin.export_log"):
        try:
            import importlib

            mod = importlib.import_module(module_name)
        except ModuleNotFoundError:
            continue
        if hasattr(mod, "SessionLocal"):
            monkeypatch.setattr(mod, "SessionLocal", testing_session_local)

    test_client = TestClient(app)
    test_client._session_local = testing_session_local  # per le assert dirette a DB
    return test_client


def _registra(client, email, nome="Utente"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _registra_admin(client, email=ADMIN_EMAIL, nome="Admin Export Lead"):
    return _registra(client, email, nome)


def _registra_consulente(client, email="consulente.export@example.com", nome="Consulente Export"):
    """Registra un utente normale e ne alza il ruolo direttamente a DB
    (stesso pattern diretto di `_seed_lead` in `test_admin_lead.py`): il
    login resta la sessione appena creata dalla registrazione, il ruolo
    aggiornato viene riletto fresco da ogni `get_current_user`."""
    resp = _registra(client, email, nome)
    with client._session_local() as session:
        user = session.query(User).filter_by(email=email).one()
        user.role = "consulente"
        session.commit()
    return resp


def _registra_user(client, email="utente.export.lead@example.com", nome="Utente Normale"):
    return _registra(client, email, nome)


def _seed_lead(session_factory, *, nome="Mario Rossi", email="mario.rossi@example.com") -> int:
    with session_factory() as session:
        lead = Lead(
            user_id=1,
            nome=nome,
            email=email,
            telefono="3331234567",
            preferenza_contatto="email",
            fascia_oraria="mattina",
            messaggio="Vorrei un sopralluogo.",
            citta="Reggio Emilia",
            mq=85,
            tipo="totale",
            totale_min=Decimal("30120.00"),
            totale_max=Decimal("44350.00"),
        )
        session.add(lead)
        session.commit()
        session.refresh(lead)
        return lead.id


class TestEsportaLead:
    def test_export_200_admin(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local)

        response = client.get(LEAD_EXPORT_URL)

        assert response.status_code == 200
        assert response.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert 'attachment; filename="lead_' in response.headers["content-disposition"]

    def test_export_200_consulente(self, client):
        """D-14: il consulente scarica i lead, non solo l'admin."""
        _registra_consulente(client)
        _seed_lead(client._session_local)

        response = client.get(LEAD_EXPORT_URL)

        assert response.status_code == 200

    def test_export_403_per_user(self, client):
        _registra_user(client)
        _seed_lead(client._session_local)

        response = client.get(LEAD_EXPORT_URL)

        assert response.status_code == 403
        with client._session_local() as db:
            assert db.query(ExportLog).count() == 0

    def test_no_store_header(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local)

        response = client.get(LEAD_EXPORT_URL)

        assert response.headers["cache-control"] == "no-store"

    def test_header_diciassette_colonne(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local)

        response = client.get(LEAD_EXPORT_URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))

        assert wb.sheetnames[0] == "ISTRUZIONI"
        istruzioni = wb["ISTRUZIONI"]
        timestamp_cell = istruzioni["B4"].value
        datetime.fromisoformat(timestamp_cell)  # non solleva -> parsabile

        assert "lead" in wb.sheetnames
        lead_sheet = wb["lead"]
        header = [cell.value for cell in next(lead_sheet.iter_rows(min_row=1, max_row=1))]
        assert header == HEADER_ATTESO
        assert len(header) == 17

    def test_righe_coincide_con_db(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local, nome="Mario Rossi")
        _seed_lead(client._session_local, nome="Luca Bianchi", email="luca.bianchi@example.com")

        response = client.get(LEAD_EXPORT_URL)

        with client._session_local() as db:
            n_righe_db = db.query(Lead).count()
            log = db.query(ExportLog).one()
            assert log.risorsa == "lead"
            assert log.include_margine is False
            assert log.righe == n_righe_db

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        n_righe_file = wb["lead"].max_row - 1
        assert n_righe_file == n_righe_db
        assert n_righe_file > 0


class TestExportLog:
    def test_lista_export_log_403_per_user(self, client):
        _registra_user(client)

        response = client.get(EXPORT_LOG_URL)

        assert response.status_code == 403

    def test_lista_export_log_ordine_decrescente_e_limite_50(self, client):
        _registra_admin(client)

        with client._session_local() as session:
            admin = session.query(User).filter_by(email=ADMIN_EMAIL).one()
            for i in range(60):
                session.add(
                    ExportLog(user_id=admin.id, risorsa="lead", righe=i, include_margine=False)
                )
            session.commit()

        response = client.get(EXPORT_LOG_URL)

        assert response.status_code == 200
        body = response.json()
        assert len(body) == 50

        created_ats = [row["created_at"] for row in body]
        assert created_ats == sorted(created_ats, reverse=True)

    def test_user_label_non_e_id_nudo(self, client):
        _registra_admin(client)
        _seed_lead(client._session_local)
        client.get(LEAD_EXPORT_URL)

        response = client.get(EXPORT_LOG_URL)
        body = response.json()
        assert len(body) == 1
        assert body[0]["user_label"] != str(body[0]["user_id"])
        assert body[0]["user_label"] == "Admin Export Lead"
        assert body[0]["risorsa"] == "lead"
