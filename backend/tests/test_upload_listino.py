"""E2E tests (RED->GREEN, TDD) per /api/v1/admin/upload-listino (ADM-01,
04-05-PLAN.md): flusso Excel a due passi preview/confirm con destinazione
(listino esistente O nuova città, D-13/D-18/D-19/D-20/D-24).

Copre: ramo listino esistente (nuove/modificate corrette), ramo nuova
città (confirm crea Listino+ComuneListino, città subito risolvibile da
/api/v1/stima senza fallback — la parte Free base_eur_mq/coefficiente resta
un ritocco CRUD separato, coerente con 04-02), riga malformata -> scarto
elencato senza bloccare il resto, `test_upload_does_not_zero_existing_rows`
(D-24), idempotenza (confirm due volte = stesso stato), validazione
dimensione/estensione (D-20), preview del ramo nuova città NON crea nulla
(T-04-42), smoke multipart reale su /preview (Pitfall 3), 403 non-admin su
preview+confirm (T-04-22).
"""

from __future__ import annotations

import io
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine

import app.db as db_module
from app.config import settings
from app.db import Base
from app.limiter import limiter
from app.main import app
from scripts.seed import load_seed_data, seed
from tests.session_factory import sessione_di_test

REGISTRAZIONE_URL = "/api/v1/auth/registrati"
STIMA_URL = "/api/v1/stima"
LISTINI_URL = "/api/v1/admin/listini"
UPLOAD_URL = "/api/v1/admin/upload-listino"

ADMIN_EMAIL = "admin.upload@example.com"


@pytest.fixture()
def client(monkeypatch, tmp_path):
    limiter.reset()
    monkeypatch.setattr(settings, "ADMIN_EMAIL", ADMIN_EMAIL)

    db_path = tmp_path / "test_upload_listino.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    testing_session_local = sessione_di_test(engine)

    with testing_session_local() as s:
        seed(s, load_seed_data())

    monkeypatch.setattr(db_module, "SessionLocal", testing_session_local)

    import app.dependencies.auth as deps_auth_module
    import app.routers.auth as auth_router_module
    import app.routers.stima as stima_router_module

    monkeypatch.setattr(deps_auth_module, "SessionLocal", testing_session_local)
    monkeypatch.setattr(auth_router_module, "SessionLocal", testing_session_local)
    monkeypatch.setattr(stima_router_module, "SessionLocal", testing_session_local)

    return TestClient(app)


def _registra_admin(client, email=ADMIN_EMAIL, nome="Admin Upload"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _registra_user(client, email="utente.normale.upload@example.com", nome="Utente Normale"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _re_listino_id(client) -> int:
    rows = client.get(LISTINI_URL).json()
    return next(r["id"] for r in rows if r["citta_riferimento"] == "Reggio Emilia")


_TOTALE_ALLESTIMENTO_ROW = (None, "TOTALE ALLESTIMENTO DI CANTIERE", None, None, None)


def _build_xlsx(rows_a_edile: list[tuple], rows_finiture: list[tuple] | None = None) -> bytes:
    """Costruisce un .xlsx minimo con i fogli A-EDILE/B-FINITURE/C-.../D-...
    (le colonne mancanti/i fogli vuoti sono gestiti come scarti a livello
    'foglio mancante' dal service, non testato qui a meno che serva).

    `rows_a_edile`: lista di tuple (colB codice, colC descrizione, colD um,
    colE qta, colF prezzo) — replica esattamente il layout reale
    (colonna indice 1..5, offset di `computo_parsing.extract_a_edile`). La
    riga di confine "TOTALE ALLESTIMENTO DI CANTIERE" viene SEMPRE
    anteposta automaticamente, cosi' le righe fornite ricadono nella fase
    di gerarchia vera (A.n/A.n.m) e non nel blocco sintetico A.0.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A-EDILE"
    for row in [_TOTALE_ALLESTIMENTO_ROW, *rows_a_edile]:
        # colonna A (indice 0) vuota, poi B..F = row
        ws.append((None, *row))

    ws_b = wb.create_sheet("B-FINITURE")
    for row in rows_finiture or []:
        ws_b.append(row)

    wb.create_sheet("C- SERRAMENTI PORTE  ")
    wb.create_sheet("D- PRODOTTI")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


A_EDILE_MINIMO = [
    ("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None),
    ("A.1.1", "Demolizione muratura", "mq", 0, 25.50),
]


def _upload_target_esistente(listino_id: int) -> str:
    return json.dumps({"kind": "esistente", "listino_id": listino_id})


def _upload_target_nuova_citta(citta: str, is_default: bool = False) -> str:
    return json.dumps({"kind": "nuova_citta", "citta": citta, "is_default": is_default})


def _post_preview(client, xlsx_bytes: bytes, target_json: str):
    return client.post(
        f"{UPLOAD_URL}/preview",
        files={"file": ("computo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"target": target_json},
    )


class TestAdminGuard403:
    """T-04-22: non-admin riceve 403 su preview+confirm."""

    def test_preview_403_per_non_admin(self, client):
        _registra_user(client)
        xlsx = _build_xlsx(A_EDILE_MINIMO)
        response = _post_preview(client, xlsx, _upload_target_nuova_citta("Carpi"))
        assert response.status_code == 403

    def test_confirm_403_per_non_admin(self, client):
        _registra_user(client)
        response = client.post(
            f"{UPLOAD_URL}/confirm",
            json={"target": {"kind": "nuova_citta", "citta": "Carpi", "is_default": False}, "nuove": [], "modificate": []},
        )
        assert response.status_code == 403

    def test_preview_401_senza_sessione(self, client):
        xlsx = _build_xlsx(A_EDILE_MINIMO)
        response = _post_preview(client, xlsx, _upload_target_nuova_citta("Carpi"))
        assert response.status_code == 401


class TestSmokeMultipart:
    """Pitfall 3: smoke test multipart reale (non solo unit-test della
    funzione di parsing) — prova che python-multipart e' correttamente
    installato/pinnato e il round-trip UploadFile+Form funziona davvero."""

    def test_preview_smoke_multipart_reale(self, client):
        _registra_admin(client)
        xlsx = _build_xlsx(A_EDILE_MINIMO)
        response = _post_preview(client, xlsx, _upload_target_nuova_citta("Guastalla"))
        assert response.status_code == 200
        body = response.json()
        assert "nuove" in body and "modificate" in body and "scarti" in body


class TestValidazioneUpload:
    """D-20: solo .xlsx, <= 10MB, rifiutato PRIMA del parsing."""

    def test_estensione_non_xlsx_respinta(self, client):
        _registra_admin(client)
        response = client.post(
            f"{UPLOAD_URL}/preview",
            files={"file": ("computo.txt", b"non e' un excel", "text/plain")},
            data={"target": _upload_target_nuova_citta("Carpi")},
        )
        assert response.status_code == 400

    def test_file_oversize_respinto(self, client):
        _registra_admin(client)
        oversize = b"0" * (10 * 1024 * 1024 + 1)
        response = client.post(
            f"{UPLOAD_URL}/preview",
            files={"file": ("computo.xlsx", oversize, "application/octet-stream")},
            data={"target": _upload_target_nuova_citta("Carpi")},
        )
        assert response.status_code == 400

    def test_destinazione_malformata_respinta(self, client):
        _registra_admin(client)
        xlsx = _build_xlsx(A_EDILE_MINIMO)
        response = _post_preview(client, xlsx, "non e' json valido")
        assert response.status_code == 400


class TestRamoListinoEsistente:
    def test_preview_righe_nuove_su_listino_esistente(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)
        xlsx = _build_xlsx(
            [("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None), ("A.1.99", "Voce mai vista prima", "mq", 0, 42.00)]
        )
        response = _post_preview(client, xlsx, _upload_target_esistente(listino_id))
        assert response.status_code == 200
        body = response.json()
        assert any(r["riga"]["descrizione"] == "Voce mai vista prima" for r in body["nuove"])
        assert body["modificate"] == []

    def test_preview_non_scrive_nulla(self, client):
        """Conteggio Voce invariato dopo /preview (niente scritture)."""
        _registra_admin(client)
        listino_id = _re_listino_id(client)
        from app.models.dettaglio_pro import Voce

        session_local = db_module.SessionLocal
        with session_local() as session:
            count_before = session.query(Voce).count()

        xlsx = _build_xlsx(
            [("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None), ("A.1.99", "Voce mai vista prima", "mq", 0, 42.00)]
        )
        _post_preview(client, xlsx, _upload_target_esistente(listino_id))

        with session_local() as session:
            count_after = session.query(Voce).count()
        assert count_after == count_before

    def test_confirm_upsert_riga_nuova_e_modificata(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)

        xlsx = _build_xlsx(
            [("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None), ("A.1.100", "Voce upload nuova", "mq", 0, 33.00)]
        )
        preview = _post_preview(client, xlsx, _upload_target_esistente(listino_id)).json()
        assert len(preview["nuove"]) == 1

        confirm = client.post(
            f"{UPLOAD_URL}/confirm",
            json={
                "target": {"kind": "esistente", "listino_id": listino_id},
                "nuove": [r["riga"] for r in preview["nuove"]],
                "modificate": [],
            },
        )
        assert confirm.status_code == 200
        body = confirm.json()
        assert body["righe_importate"] == 1
        assert body["listino_id"] == listino_id

        # Ri-uploadare la stessa riga con prezzo diverso -> modificata, non duplicata
        xlsx2 = _build_xlsx(
            [("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None), ("A.1.100", "Voce upload nuova", "mq", 0, 55.00)]
        )
        preview2 = _post_preview(client, xlsx2, _upload_target_esistente(listino_id)).json()
        assert preview2["nuove"] == []
        assert len(preview2["modificate"]) == 1
        assert float(preview2["modificate"][0]["valore_precedente"]) == 33.00

    def test_riga_malformata_scartata_resto_importato(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)
        xlsx = _build_xlsx(
            [
                ("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None),
                ("A.1.101", "Voce con prezzo zero", "mq", 0, 0),
                ("A.1.102", "Voce valida", "mq", 0, 20.00),
            ]
        )
        preview = _post_preview(client, xlsx, _upload_target_esistente(listino_id)).json()
        assert len(preview["scarti"]) >= 1
        assert any(r["riga"]["descrizione"] == "Voce valida" for r in preview["nuove"])
        assert not any(r["riga"]["descrizione"] == "Voce con prezzo zero" for r in preview["nuove"])

    def test_upload_does_not_zero_existing_rows(self, client):
        """D-24: upload additivo, MAI un DELETE implicito di righe assenti
        dal file caricato."""
        _registra_admin(client)
        listino_id = _re_listino_id(client)
        from app.models.dettaglio_pro import Voce

        session_local = db_module.SessionLocal
        with session_local() as session:
            count_before = session.query(Voce).count()
            assert count_before > 0, "seed atteso: gia' voci presenti su RE"

        # Upload con UNA SOLA voce (un file "parziale" rispetto al DB reale)
        xlsx = _build_xlsx(
            [("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None), ("A.1.103", "Voce isolata", "mq", 0, 10.00)]
        )
        preview = _post_preview(client, xlsx, _upload_target_esistente(listino_id)).json()
        client.post(
            f"{UPLOAD_URL}/confirm",
            json={
                "target": {"kind": "esistente", "listino_id": listino_id},
                "nuove": [r["riga"] for r in preview["nuove"]],
                "modificate": [r["riga"] for r in preview["modificate"]],
            },
        )

        with session_local() as session:
            count_after = session.query(Voce).count()
        # Le voci precedenti restano TUTTE, piu' la nuova aggiunta
        assert count_after == count_before + 1

    def test_confirm_idempotente_stesso_stato(self, client):
        _registra_admin(client)
        listino_id = _re_listino_id(client)
        xlsx = _build_xlsx(
            [("A.1", "DEMOLIZIONI - RIMOZIONI", None, None, None), ("A.1.104", "Voce idempotente", "mq", 0, 15.00)]
        )
        preview = _post_preview(client, xlsx, _upload_target_esistente(listino_id)).json()
        payload = {
            "target": {"kind": "esistente", "listino_id": listino_id},
            "nuove": [r["riga"] for r in preview["nuove"]],
            "modificate": [],
        }
        first = client.post(f"{UPLOAD_URL}/confirm", json=payload)
        assert first.status_code == 200

        from app.models.dettaglio_pro import Voce

        session_local = db_module.SessionLocal
        with session_local() as session:
            count_after_first = session.query(Voce).filter_by(descrizione="Voce idempotente").count()
        assert count_after_first == 1

        # Ri-confermare lo stesso payload (client re-invia lo stesso diff) non duplica
        second = client.post(f"{UPLOAD_URL}/confirm", json=payload)
        assert second.status_code == 200
        with session_local() as session:
            count_after_second = session.query(Voce).filter_by(descrizione="Voce idempotente").count()
        assert count_after_second == 1

    def test_confirm_listino_inesistente_404(self, client):
        _registra_admin(client)
        response = client.post(
            f"{UPLOAD_URL}/confirm",
            json={"target": {"kind": "esistente", "listino_id": 999999}, "nuove": [], "modificate": []},
        )
        assert response.status_code == 404


class TestRamoNuovaCitta:
    """D-13/T-04-42: preview non crea nulla, confirm crea Listino+ComuneListino."""

    def test_preview_nuova_citta_non_crea_listino(self, client):
        _registra_admin(client)
        listini_before = client.get(LISTINI_URL).json()
        count_before = len(listini_before)

        xlsx = _build_xlsx(A_EDILE_MINIMO)
        response = _post_preview(client, xlsx, _upload_target_nuova_citta("Rubiera"))
        assert response.status_code == 200
        body = response.json()
        assert len(body["nuove"]) >= 1  # tutte le righe sono 'nuove' (stato di confronto vuoto)
        assert body["modificate"] == []

        listini_after = client.get(LISTINI_URL).json()
        assert len(listini_after) == count_before
        assert not any(row["citta_riferimento"] == "Rubiera" for row in listini_after)

    def test_confirm_nuova_citta_crea_listino_e_comune_listino(self, client):
        _registra_admin(client)
        xlsx = _build_xlsx(A_EDILE_MINIMO)
        preview = _post_preview(client, xlsx, _upload_target_nuova_citta("Correggio")).json()

        confirm = client.post(
            f"{UPLOAD_URL}/confirm",
            json={
                "target": {"kind": "nuova_citta", "citta": "Correggio", "is_default": False},
                "nuove": [r["riga"] for r in preview["nuove"]],
                "modificate": [],
            },
        )
        assert confirm.status_code == 200
        body = confirm.json()
        assert body["righe_importate"] >= 1
        new_listino_id = body["listino_id"]

        listini = client.get(LISTINI_URL).json()
        created = next(row for row in listini if row["id"] == new_listino_id)
        assert created["citta_riferimento"] == "Correggio"

        from app.models.catalogo import ComuneListino

        session_local = db_module.SessionLocal
        with session_local() as session:
            comune = session.query(ComuneListino).filter_by(comune_slug="correggio").one_or_none()
            assert comune is not None
            assert comune.listino_id == new_listino_id

    def test_confirm_nuova_citta_risolvibile_senza_fallback(self, client):
        """D-13: dopo confirm, `resolve_listino` risolve la nuova città senza
        fallback (ComuneListino creato). Il completamento della catena Free
        (`base_eur_mq`/`coefficiente_impresa` su ListinoCategoria) resta un
        ritocco CRUD separato e deliberato (04-02), fuori scope per questo
        upload che tocca solo Sottocategoria/Voce (dettaglio Pro, D-10)."""
        _registra_admin(client)
        xlsx = _build_xlsx(A_EDILE_MINIMO)
        preview = _post_preview(client, xlsx, _upload_target_nuova_citta("Scandiano")).json()
        client.post(
            f"{UPLOAD_URL}/confirm",
            json={
                "target": {"kind": "nuova_citta", "citta": "Scandiano", "is_default": False},
                "nuove": [r["riga"] for r in preview["nuove"]],
                "modificate": [],
            },
        )

        from app.services.city_resolver import resolve_listino

        session_local = db_module.SessionLocal
        with session_local() as session:
            listino, fallback = resolve_listino(session, "Scandiano")
            assert fallback is False
            assert listino.citta_riferimento == "Scandiano"

    def test_confirm_nuova_citta_idempotente_non_duplica_listino(self, client):
        _registra_admin(client)
        xlsx = _build_xlsx(A_EDILE_MINIMO)
        preview = _post_preview(client, xlsx, _upload_target_nuova_citta("Novellara")).json()
        payload = {
            "target": {"kind": "nuova_citta", "citta": "Novellara", "is_default": False},
            "nuove": [r["riga"] for r in preview["nuove"]],
            "modificate": [],
        }
        first = client.post(f"{UPLOAD_URL}/confirm", json=payload).json()
        second = client.post(f"{UPLOAD_URL}/confirm", json=payload).json()
        assert first["listino_id"] == second["listino_id"]

        listini = client.get(LISTINI_URL).json()
        novellara_rows = [row for row in listini if row["citta_riferimento"] == "Novellara"]
        assert len(novellara_rows) == 1
