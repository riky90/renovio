"""Test dell'export del catalogo domande (`GET /admin/domande/export.xlsx`,
D-05, D-06, D-07, D-15). Come `test_export_margine.py`: apre il file
generato con openpyxl e guarda dentro, perche' e' il contenuto che viaggia
fuori dall'applicazione (regola di dominio #1, `CLAUDE.md`).

Copre: cinque fogli dati con header `id`-first e conteggio uguale alla
tabella corrispondente; la colonna `regola` del foglio `derivazioni` resa
leggibile; l'audit `export_log`; il 403 per consulente/utente senza nessuna
riga di audit; nessuna colonna `impresa`/`coefficiente` in nessun foglio.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine

import app.db as db_module
from app.config import settings
from app.db import Base
from app.limiter import limiter
from app.main import app
from app.models.domanda import (
    AttesaImmobile,
    DerivazioneRisposta,
    Domanda,
    EffettoOpzione,
    OpzioneDomanda,
)
from app.models.export_log import ExportLog
from app.models.user import User
from app.services.export_excel.domande import (
    COLONNE_ATTESE,
    COLONNE_DERIVAZIONI,
    COLONNE_DOMANDE,
    COLONNE_EFFETTI,
    COLONNE_OPZIONI,
)
from scripts.seed import load_seed_data, seed
from tests.session_factory import sessione_di_test

REGISTRAZIONE_URL = "/api/v1/auth/registrati"
URL = "/api/v1/admin/domande/export.xlsx"

ADMIN_EMAIL = "admin.export.domande@example.com"

FOGLI_DATI = {
    "domande": (COLONNE_DOMANDE, Domanda),
    "opzioni": (COLONNE_OPZIONI, OpzioneDomanda),
    "effetti": (COLONNE_EFFETTI, EffettoOpzione),
    "attese": (COLONNE_ATTESE, AttesaImmobile),
    "derivazioni": (COLONNE_DERIVAZIONI, DerivazioneRisposta),
}


@pytest.fixture()
def client(monkeypatch, tmp_path):
    limiter.reset()
    monkeypatch.setattr(settings, "ADMIN_EMAIL", ADMIN_EMAIL)

    db_path = tmp_path / "test_export_domande.db"
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    testing_session_local = sessione_di_test(engine)

    with testing_session_local() as s:
        seed(s, load_seed_data())

    monkeypatch.setattr(db_module, "SessionLocal", testing_session_local)

    import app.dependencies.auth as deps_auth_module
    import app.routers.auth as auth_router_module

    monkeypatch.setattr(deps_auth_module, "SessionLocal", testing_session_local)
    monkeypatch.setattr(auth_router_module, "SessionLocal", testing_session_local)

    test_client = TestClient(app)
    test_client._session_local = testing_session_local
    return test_client


def _registra_admin(client, email=ADMIN_EMAIL, nome="Admin Export Domande"):
    return client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})


def _registra_ruolo(client, *, role, email, nome):
    resp = client.post(REGISTRAZIONE_URL, json={"email": email, "password": "Password123", "nome": nome})
    with client._session_local() as db:
        user = db.query(User).filter_by(email=email).one()
        user.role = role
        db.commit()
    return resp


class TestExportOk:
    def test_status_200_content_type_e_filename(self, client):
        _registra_admin(client)

        response = client.get(URL)

        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "catalogo-domande" in response.headers["content-disposition"]

    def test_cinque_fogli_piu_istruzioni_nessun_avviso(self, client):
        _registra_admin(client)

        response = client.get(URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))

        attesi = {"ISTRUZIONI", "domande", "opzioni", "effetti", "attese", "derivazioni"}
        assert attesi.issubset(set(wb.sheetnames))
        assert "AVVISO" not in wb.sheetnames

    @pytest.mark.parametrize("nome_foglio", list(FOGLI_DATI))
    def test_header_id_prima_colonna_e_uguale_a_colonne(self, client, nome_foglio):
        _registra_admin(client)

        response = client.get(URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        colonne_attese, _ = FOGLI_DATI[nome_foglio]

        ws = wb[nome_foglio]
        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

        assert header[0] == "id"
        assert header == colonne_attese

    @pytest.mark.parametrize("nome_foglio", list(FOGLI_DATI))
    def test_conteggio_righe_uguale_alla_tabella_e_maggiore_di_zero(self, client, nome_foglio):
        _registra_admin(client)

        response = client.get(URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        _, modello = FOGLI_DATI[nome_foglio]

        ws = wb[nome_foglio]
        n_righe_file = ws.max_row - 1  # meno l'header

        with client._session_local() as db:
            n_righe_db = db.query(modello).count()

        assert n_righe_file == n_righe_db
        assert n_righe_file > 0

    def test_id_domande_contiene_i_code(self, client):
        _registra_admin(client)

        response = client.get(URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb["domande"]

        with client._session_local() as db:
            code_attesi = {d.code for d in db.query(Domanda).all()}

        id_nel_file = {row[0].value for row in ws.iter_rows(min_row=2)}
        assert id_nel_file == code_attesi

    def test_derivazioni_regola_pre_1980_contiene_i_testi_delle_due_domande(self, client):
        _registra_admin(client)

        with client._session_local() as db:
            derivazione = (
                db.query(DerivazioneRisposta).filter_by(da_opzione_code="pre_1980").one()
            )
            testo_da_domanda = db.get(Domanda, derivazione.da_domanda_code).testo
            testo_domanda = db.get(Domanda, derivazione.domanda_code).testo

        response = client.get(URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb["derivazioni"]

        colonne = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        idx_id = colonne.index("id")
        idx_regola = colonne.index("regola")

        riga = next(
            row for row in ws.iter_rows(min_row=2) if row[idx_id].value == derivazione.id
        )
        regola = riga[idx_regola].value

        assert testo_da_domanda in regola
        assert testo_domanda in regola

    def test_export_log_scritto_con_somma_delle_righe(self, client):
        _registra_admin(client)

        response = client.get(URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        somma = sum(wb[nome].max_row - 1 for nome in FOGLI_DATI)

        with client._session_local() as db:
            righe = db.query(ExportLog).filter_by(risorsa="domande").all()
            assert len(righe) == 1
            assert righe[0].include_margine is False
            assert righe[0].righe == somma

    def test_nessuna_colonna_impresa_o_coefficiente(self, client):
        _registra_admin(client)

        response = client.get(URL)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))

        for nome_foglio in wb.sheetnames:
            ws = wb[nome_foglio]
            header_row = next(ws.iter_rows(min_row=1, max_row=1), None)
            if header_row is None:
                continue
            header = [str(cell.value) for cell in header_row if cell.value is not None]
            for cella in header:
                assert "impresa" not in cella.lower()
                assert "coefficiente" not in cella.lower()


class TestExportVietatoAiNonAdmin:
    def test_403_per_consulente_e_utente_normale_senza_audit(self, client):
        _registra_ruolo(
            client, role="consulente", email="consulente.export.domande@example.com", nome="Consulente"
        )
        response_consulente = client.get(URL)
        assert response_consulente.status_code == 403

        _registra_ruolo(client, role="user", email="utente.export.domande@example.com", nome="Utente")
        response_utente = client.get(URL)
        assert response_utente.status_code == 403

        with client._session_local() as db:
            assert db.query(ExportLog).count() == 0
