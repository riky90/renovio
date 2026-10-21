"""Computo parametrico (Fase 11): grandezze, regole per voce, preset, motore
unico Free/Pro.

Cosa va dimostrato:

  1. IL COMPUTO SEMINATO E' COERENTE: nessun errore di curatela e ogni voce
     del listino di default ha la sua regola (una voce senza regola ricadrebbe
     sulla catena v2, cioe' sul bug delle quantita' x m2).
  2. FREE E PRO SONO LO STESSO NUMERO (MOT-01): per ogni categoria il Free e'
     il Pro piu' IVA, a meno degli arrotondamenti all'euro delle voci.
  3. LE VOCI IN PERCENTUALE VALGONO (RGL-04): assistenze e smaltimento sono
     calcolate dalle loro sorgenti, mai 0 per costruzione.
  4. LE ALTERNATIVE NON SI SOMMANO e le voci che il tecnico toglierebbe non
     entrano.
  5. IL TIPO E' UN PRESET (INT-02): sceglie cosa si fa; una risposta esplicita
     vince sul preset.
  6. UN ERRORE DI CURATELA NON FA 500: un ciclo vale 0 e la validazione lo dice.
  7. UN LISTINO SENZA REGOLE RICADE SUL CALCOLO EUR/m2 (MOT-03): l'oracolo v1
     resta vero li'.
  8. L'EXCEL FA IL GIRO COMPLETO ed e' tutto o niente.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import create_engine

from app.db import Base
from app.models.catalogo import Categoria
from app.models.computo import Grandezza, RegolaVoce, SorgenteRegola, TermineGrandezza
from app.models.dettaglio_pro import Sottocategoria, Voce
from app.services.city_resolver import resolve_listino
from app.services.computo import computo_free, risolvi_grandezze, valida_computo
from app.services.computo_excel import esporta_computo, importa_computo
from app.services.dettaglio_pro import DettaglioProService
from app.services.estimation import EstimationService
from scripts.seed import load_seed_data, seed
from tests.session_factory import sessione_di_test


@pytest.fixture()
def session():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    with sessione_di_test(engine)() as s:
        seed(s, load_seed_data())
        yield s


@pytest.fixture()
def listino(session):
    listino, fallback = resolve_listino(session, "Reggio Emilia")
    assert fallback is False
    return listino


def _voci(session, listino, mq=85, tipo="totale", risposte=None):
    risultato = computo_free(session, listino, mq=mq, tipo=tipo, risposte=risposte or {})
    assert risultato is not None
    return risultato, {v.codice: v for v in risultato.voci}


# --- 1. coerenza del seed ---------------------------------------------------


def test_computo_seminato_senza_errori_di_curatela(session):
    assert valida_computo(session) == []


def test_ogni_voce_del_listino_di_default_ha_una_regola(session, listino):
    codici = {
        v.codice
        for v in session.query(Voce)
        .join(Sottocategoria, Sottocategoria.id == Voce.sottocategoria_id)
        .filter(Sottocategoria.listino_id == listino.id)
        .all()
    }
    assert None not in codici, "una voce senza codice non puo' avere una regola"
    regole = {r.voce_codice for r in session.query(RegolaVoce).all()}
    assert codici - regole == set()


# --- 2. un motore solo -------------------------------------------------------


@pytest.mark.parametrize("mq", [45, 85, 140])
@pytest.mark.parametrize("tipo", ["leggera", "moderata", "totale"])
def test_free_e_pro_coincidono_per_categoria(session, listino, mq, tipo):
    free = EstimationService.compute(session, listino, mq=mq, tipo=tipo, risposte={})
    pro = DettaglioProService.compute(session, listino, mq, tipo, {})
    iva = {c.code: c.iva_pct for c in session.query(Categoria).all()}
    pro_per_cat = {c.code: (c.totale_min, c.totale_max) for c in pro.categorie}
    for cat in free.categorie:
        pro_min, pro_max = pro_per_cat.get(cat.code, (Decimal("0"), Decimal("0")))
        n_voci = sum(len(s.voci) for c in pro.categorie if c.code == cat.code for s in c.sottocategorie)
        tolleranza = (Decimal(n_voci) + 2) * (1 + iva[cat.code])
        assert abs(cat.totale_min - pro_min * (1 + iva[cat.code])) <= tolleranza, cat.code
        assert abs(cat.totale_max - pro_max * (1 + iva[cat.code])) <= tolleranza, cat.code


def test_risposte_none_vale_come_preset_del_tipo(session, listino):
    senza = EstimationService.compute(session, listino, mq=85, tipo="moderata", risposte=None)
    vuote = EstimationService.compute(session, listino, mq=85, tipo="moderata", risposte={})
    assert (senza.totale_min, senza.totale_max) == (vuote.totale_min, vuote.totale_max)


def test_totale_85mq_nell_ordine_di_grandezza_del_computo(session, listino):
    """Il Pro v2 dava 386-579 k EUR su questo caso; ricostruito a mano col
    file, l'edile imponibile vale circa 42 k EUR. Il numero esatto e'
    bozza (TAR-01): qui si fissa solo l'ordine di grandezza."""
    risultato, _ = _voci(session, listino)
    edile = risultato.per_categoria["A-EDILE"]
    assert Decimal("25000") < edile < Decimal("60000")
    assert sum(risultato.per_categoria.values()) < Decimal("120000")


# --- 3. voci in percentuale ---------------------------------------------------


def test_smaltimento_e_il_venti_per_cento_delle_demolizioni(session, listino):
    risultato, voci = _voci(session, listino)
    demolizioni = sum(
        v.importo
        for v in risultato.voci
        if v.codice and v.codice.startswith("A.1.") and v.codice.count(".") >= 2
    )
    assert demolizioni > 0
    assert voci["A.9.6"].importo == pytest.approx(demolizioni * Decimal("0.20"))


def test_assistenze_murarie_non_valgono_zero(session, listino):
    _, voci = _voci(session, listino)
    assert voci["A.9.1"].importo > 0
    assert voci["A.9.2"].importo > 0
    assert voci["A.0.18"].importo > 0, "movimentazione: 6% dei pavimenti x piano"


def test_assistenza_impianti_e_il_trenta_per_cento(session, listino):
    _, voci = _voci(session, listino)
    sorgenti = [s.sorgente_codice for s in session.query(SorgenteRegola).filter_by(voce_codice="A.9.1")]
    base = sum((voci[c].importo for c in sorgenti if c in voci), Decimal("0"))
    assert voci["A.9.1"].importo == pytest.approx(base * Decimal("0.30"))


# --- 4. alternative ---------------------------------------------------------


@pytest.mark.parametrize("tipo", ["leggera", "moderata", "totale"])
def test_alternative_mai_sommate(session, listino, tipo):
    _, voci = _voci(session, listino, tipo=tipo)
    gruppi = [
        {"A.2.1", "A.2.2", "A.2.3"},  # massetti
        {"A.3.1", "A.3.2", "A.3.3"},  # murature nuove
        {"A.8.15", "A.8.16"},  # elettrico modifica / rifacimento
        {"A.8.10", "A.8.13.1"},  # radiatori / pavimento radiante
        {"A.8.21", "A.8.22"},  # split / canalizzato
        {"C.1.1", "C.1.3", "C.1.4"},  # porte battente / scorrevoli
        {"C.1.11", "C.1.12"},  # tapparelle PVC / alluminio
        {"A.1.5", "A.2.5"},  # demolisco e rifaccio / posa sopra
    ]
    for gruppo in gruppi:
        assert len(gruppo & voci.keys()) <= 1, gruppo


def test_voci_da_togliere_non_entrano(session, listino):
    _, voci = _voci(session, listino)
    for codice in ("A.1.13", "A.1.15", "A.1.16", "A.2.4"):
        assert codice not in voci


def test_voci_a_corpo_non_scalano_coi_metri_quadri(session, listino):
    """La causa n. 3 della v2: voci 'ac'/'cad' moltiplicate per i m2."""
    for mq in (60, 200):
        _, voci = _voci(session, listino, mq=mq)
        for codice in ("A.6.11", "A.9.3", "A.0.13", "A.0.15"):
            if codice in voci:
                assert voci[codice].quantita <= 10, (codice, voci[codice].quantita)


# --- 5. preset ---------------------------------------------------------------


def test_il_tipo_sceglie_cosa_si_fa(session, listino):
    _, leggera = _voci(session, listino, tipo="leggera")
    _, totale = _voci(session, listino, tipo="totale")
    assert "A.8.1" not in leggera, "una leggera non rifa' l'impianto idrico"
    assert "A.8.1" in totale
    assert "A.2.5" in leggera and "A.1.5" in totale


def test_risposta_esplicita_vince_sul_preset(session, listino):
    _, voci = _voci(session, listino, tipo="leggera", risposte={"interv_bagni": "tutti", "bagni": 2})
    assert voci["A.8.5"].quantita == 8, "quattro apparecchi per due bagni"


def test_grandezza_da_risposta_numerica(session):
    grandezze = risolvi_grandezze(session, mq=85, tipo="totale", risposte={"bagni": 3})
    assert grandezze["bagni"].valore == 3
    assert grandezze["bagni"].origine == "risposta"
    assert grandezze["bagni_rifatti"].valore == 3


def test_split_solo_con_la_climatizzazione(session, listino):
    _, senza = _voci(session, listino, risposte={"interv_clima": "nessuno"})
    assert "A.8.21" not in senza
    _, con = _voci(session, listino, risposte={"interv_clima": "split", "condizionatori": 3})
    assert con["A.8.21"].quantita == 3


# --- 6. errori di curatela ------------------------------------------------------


def test_ciclo_fra_grandezze_vale_zero_e_si_segnala(session, listino):
    session.add_all(
        [
            Grandezza(code="ciclo_a", nome="A"),
            Grandezza(code="ciclo_b", nome="B"),
        ]
    )
    session.flush()
    session.add_all(
        [
            TermineGrandezza(grandezza_code="ciclo_a", sorgente_code="ciclo_b", fattore=Decimal("1")),
            TermineGrandezza(grandezza_code="ciclo_b", sorgente_code="ciclo_a", fattore=Decimal("1")),
        ]
    )
    session.flush()
    grandezze = risolvi_grandezze(session, mq=85, tipo="totale", risposte={})
    assert grandezze["ciclo_a"].valore == 0
    assert any("Ciclo fra grandezze" in e for e in valida_computo(session))


def test_ciclo_fra_voci_non_solleva(session, listino):
    session.add(SorgenteRegola(voce_codice="A.2.2", sorgente_codice="A.2.2"))
    regola = session.get(RegolaVoce, "A.4.2")
    regola.tipo_regola = "da_voci"
    session.add(SorgenteRegola(voce_codice="A.4.2", sorgente_codice="A.11.2"))
    regola_rasatura = session.get(RegolaVoce, "A.11.2")
    regola_rasatura.tipo_regola = "da_voci"
    session.add(SorgenteRegola(voce_codice="A.11.2", sorgente_codice="A.4.2"))
    session.flush()
    _voci(session, listino)
    assert any("Ciclo fra voci" in e for e in valida_computo(session))


def test_voce_a_corpo_proporzionale_ai_mq_e_un_errore(session):
    regola = session.get(RegolaVoce, "A.6.11")
    regola.tipo_regola = "grandezza"
    regola.grandezza_code = "mq"
    session.flush()
    assert any("A.6.11" in e and "a corpo" in e for e in valida_computo(session))


# --- 7. fallback ------------------------------------------------------------------


def test_listino_senza_regole_torna_al_calcolo_eur_mq(session, listino):
    for regola in session.query(RegolaVoce).all():
        regola.attiva = False
    session.flush()
    assert computo_free(session, listino, mq=85, tipo="totale", risposte={}) is None
    result = EstimationService.compute(session, listino, mq=85, tipo="totale")
    assert (result.totale_min, result.totale_max) == (Decimal("47985"), Decimal("70575"))


# --- 8. Excel -------------------------------------------------------------------------


def _workbook(file_bytes: bytes):
    return openpyxl.load_workbook(io.BytesIO(file_bytes))


def test_excel_giro_completo_senza_differenze(session, listino):
    prima = computo_free(session, listino, mq=85, tipo="totale", risposte={})
    file_bytes, righe = esporta_computo(session)
    assert righe > 0
    esito = importa_computo(session, _workbook(file_bytes), applica=True)
    assert esito.errori == []
    assert esito.applicato
    dopo = computo_free(session, listino, mq=85, tipo="totale", risposte={})
    assert dopo.per_categoria == prima.per_categoria


def test_excel_con_errori_non_cambia_nulla(session, listino):
    prima = computo_free(session, listino, mq=85, tipo="totale", risposte={})
    file_bytes, _ = esporta_computo(session)
    wb = _workbook(file_bytes)
    ws = wb["regole"]
    intestazione = [c.value for c in ws[1]]
    ws.cell(row=2, column=intestazione.index("tipo") + 1, value="formula_libera")
    esito = importa_computo(session, wb, applica=True)
    assert not esito.applicato
    assert any("formula_libera" in e for e in esito.errori)
    dopo = computo_free(session, listino, mq=85, tipo="totale", risposte={})
    assert dopo.per_categoria == prima.per_categoria


def test_excel_anteprima_non_scrive(session, listino):
    file_bytes, _ = esporta_computo(session)
    wb = _workbook(file_bytes)
    ws = wb["regole"]
    intestazione = [c.value for c in ws[1]]
    riga = next(
        r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "A.12.1"
    )
    ws.cell(row=riga, column=intestazione.index("fattore") + 1, value=3)
    esito = importa_computo(session, wb, applica=False)
    assert esito.errori == [] and not esito.applicato
    assert session.get(RegolaVoce, "A.12.1").fattore == 1


def test_excel_non_contiene_il_coefficiente_impresa(session):
    file_bytes, _ = esporta_computo(session)
    wb = _workbook(file_bytes)
    for ws in wb.worksheets:
        for riga in ws.iter_rows(values_only=True):
            for valore in riga:
                assert "impresa" not in str(valore or "").lower()
